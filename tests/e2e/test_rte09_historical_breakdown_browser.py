"""RTE09 E01: periodo histórico y desglose interactivo, en el navegador.

Lo que se comprueba
-------------------
* Los tres recorridos: Week → día, Month → semana, Year → mes; cada uno acota
  el *Activity breakdown* y **no** cambia los KPIs.
* La fila elegida se distingue sin depender del color (`aria-pressed` y la
  marca *Selected*) y *Whole …* vuelve al periodo entero.
* Cambiar supervisor, periodo concreto o pestaña borra el subgrupo.
* El Excel, con un subgrupo elegido, sigue siendo el del periodo principal.
* Una recarga conserva el periodo elegido.
* Con el reloj del navegador en 2024, la pantalla abre el periodo en curso del
  **servidor** y no manda ninguna fecha.
* Móvil a 390 px: los tres controles a la vista, sin desbordamiento.

Datos: el supervisor tiene jornadas el martes 15 (Client Visit, 10 mi, tres
actividades) y el miércoles 16 de septiembre de 2026 (Office, 3 mi), el 10 de
enero de 2026 (Office, 7 mi) y una de hoy, que hace que el periodo en curso sea
el actual. El `route_admin`, que también es supervisor, tiene una el día 15.

Capturas en `var/screenshots/rte09-e01/`, fuera del árbol versionado.
"""

from __future__ import annotations

import pathlib
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from playwright.async_api import expect

from tests.e2e.conftest import abrir_sesion, lanzar_edge
from tests.integration.test_activity_explorer import _jornada_con_parada, _perfil_y_vehiculo
from tests.integration.test_reports import _jornada

pytestmark = [pytest.mark.integration, pytest.mark.browser]

pytest.importorskip("playwright", reason="playwright no está instalado")

DESTINO = pathlib.Path("var/screenshots/rte09-e01")
RUTA = "/admin/route/reports"
ESCRITORIO = {"width": 1280, "height": 900}
MOVIL = {"width": 390, "height": 844}
KPIS = ("Total miles", "Activity time", "Activities", "Est. fuel cost")


async def _historia(alpha_client, seeded) -> None:
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-B1")
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-B2", usuario="route_admin")
    await _jornada(alpha_client, seeded, dia=date(2026, 1, 10), millas="7.0", proposito="office")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0", actividades=3)
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 16), millas="3.0", proposito="office")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="4.0", usuario="route_admin")
    # Hoy: el periodo en curso es el actual, no el de la última jornada fechada.
    await _jornada_con_parada(alpha_client, seeded)


async def _abrir(navegador, live_server, seeded, viewport, *, tactil: bool = False):
    contexto = await navegador.new_context(
        base_url=live_server, viewport=viewport, accept_downloads=True, has_touch=tactil
    )
    page = await contexto.new_page()
    await abrir_sesion(page, seeded.alpha.users["route_admin"].email)
    await page.goto(RUTA)
    await expect(page.locator("#report-period")).to_be_enabled(timeout=20_000)
    return contexto, page


def _es_informe(r, **esperado) -> bool:
    if "/api/reports" not in r.url or "/export" in r.url:
        return False
    return all(f"{k}={v}" in r.url for k, v in esperado.items())


async def _periodo(page, inicio: str) -> dict:
    """Elige un periodo concreto y devuelve la lectura que eso disparó."""
    async with page.expect_response(lambda r: _es_informe(r, date=inicio)) as info:
        await page.locator("#report-period").select_option(inicio)
    return await (await info.value).json()


async def _pestana(page, nombre: str) -> dict:
    async with page.expect_response(
        lambda r: _es_informe(r, range=nombre.lower()) and "group=" not in r.url
    ) as info:
        await page.get_by_role("button", name=nombre, exact=True).click()
    return await (await info.value).json()


async def _grupo(page, inicio: str) -> dict:
    async with page.expect_response(lambda r: _es_informe(r, group=inicio)) as info:
        await page.locator(f'[data-group-start="{inicio}"]').click()
    return await (await info.value).json()


async def _ano_2026(page) -> None:
    """En Year, 2026: el en curso mientras se ejecute en 2026; si no, se elige."""
    if await page.locator("#report-period").input_value() != "2026-01-01":
        await _periodo(page, "2026-01-01")


async def _kpis(page) -> list[str]:
    return [await page.locator(f'[data-kpi="{k}"]').inner_text() for k in KPIS]


def _alcance(page):
    return page.locator("[data-breakdown-scope]")


def _elegidas(page):
    return page.locator('[data-group-start][aria-pressed="true"]')


async def _sin_desbordar(page, cuando: str) -> None:
    ancho = await page.evaluate(
        "document.documentElement.scrollWidth - document.documentElement.clientWidth"
    )
    assert ancho <= 0, f"{cuando}: desborda {ancho}px"


async def test_escritorio_los_tres_recorridos_y_el_excel(seeded, alpha_client, live_server):
    from playwright.async_api import async_playwright

    DESTINO.mkdir(parents=True, exist_ok=True)
    await _historia(alpha_client, seeded)

    async with async_playwright() as p:
        navegador = await lanzar_edge(p)
        try:
            contexto, page = await _abrir(navegador, live_server, seeded, ESCRITORIO)
            await expect(page.locator('label[for="report-period"]')).to_have_text("Week")

            # ── Week → día ──
            semana = await _periodo(page, "2026-09-14")
            assert semana["start"] == "2026-09-14"
            await expect(page.locator("#report-period")).to_have_value("2026-09-14")
            await expect(page.get_by_text("All Supervisors · Sep 14–20, 2026")).to_be_visible()
            await expect(_alcance(page)).to_have_text("Showing whole week")
            antes = await _kpis(page)
            await expect(page.locator('[data-kpi="Total miles"]')).to_contain_text("17.0")

            dia = await _grupo(page, "2026-09-16")
            assert [r["purpose"] for r in dia["breakdown"]] == ["office"]
            await expect(_alcance(page)).to_have_text("Showing Wed Sep 16 only")
            await expect(_elegidas(page)).to_have_count(1)
            await expect(_elegidas(page)).to_contain_text("Selected")
            await expect(page.get_by_text("Client Visit", exact=True)).to_have_count(0)
            assert await _kpis(page) == antes, "los KPIs son del periodo principal"
            await page.screenshot(path=str(DESTINO / "escritorio-week-dia.png"), full_page=True)

            # El Excel con un día elegido: la semana entera.
            async with page.expect_download() as info:
                await page.get_by_role("button", name="Export Excel").click()
            descarga = await info.value
            assert "2026-09-14_2026-09-20" in descarga.suggested_filename
            hoja = load_workbook(
                BytesIO(pathlib.Path(await descarga.path()).read_bytes())
            )["Activity Detail"]
            assert len(list(hoja.iter_rows(values_only=True))) == 1 + 3

            # Whole week: vuelve al periodo entero sin otra petición.
            await page.get_by_role("button", name="Whole week").click()
            await expect(_elegidas(page)).to_have_count(0)
            await expect(_alcance(page)).to_have_text("Showing whole week")
            await expect(page.get_by_text("Client Visit", exact=True)).to_be_visible()

            # Con teclado: Tab hasta la fila y Enter.
            await page.locator('[data-group-start="2026-09-15"]').focus()
            async with page.expect_response(lambda r: _es_informe(r, group="2026-09-15")):
                await page.keyboard.press("Enter")
            await expect(_alcance(page)).to_have_text("Showing Tue Sep 15 only")

            # ── Month → semana ──
            await _pestana(page, "Month")
            await expect(_elegidas(page)).to_have_count(0)
            await expect(page.locator('label[for="report-period"]')).to_have_text("Month")
            await _periodo(page, "2026-09-01")
            await expect(page.get_by_text("By week", exact=True)).to_be_visible()
            antes = await _kpis(page)
            mes = await _grupo(page, "2026-09-14")
            assert mes["selected_group"] == {"start": "2026-09-14", "end": "2026-09-20"}
            await expect(_alcance(page)).to_have_text("Showing Sep 14–20 only")
            assert await _kpis(page) == antes
            await page.screenshot(path=str(DESTINO / "escritorio-month-semana.png"), full_page=True)

            # ── Year → mes ──
            await _pestana(page, "Year")
            await expect(_elegidas(page)).to_have_count(0)
            await expect(_alcance(page)).to_have_text("Showing whole year")
            await expect(page.locator('label[for="report-period"]')).to_have_text("Year")
            await _ano_2026(page)
            antes = await _kpis(page)
            ano = await _grupo(page, "2026-01-01")
            assert [(r["purpose"], r["official_miles"]) for r in ano["breakdown"]] == [
                ("office", "7.0")
            ]
            await expect(_alcance(page)).to_have_text("Showing January only")
            assert await _kpis(page) == antes
            await page.screenshot(path=str(DESTINO / "escritorio-year-mes.png"), full_page=True)
            await contexto.close()
        finally:
            await navegador.close()


async def test_cambiar_el_contexto_borra_el_subgrupo_y_la_recarga_lo_conserva(
    seeded, alpha_client, live_server
):
    from playwright.async_api import async_playwright

    await _historia(alpha_client, seeded)
    admin = seeded.alpha.users["route_admin"].id

    async with async_playwright() as p:
        navegador = await lanzar_edge(p)
        try:
            contexto, page = await _abrir(navegador, live_server, seeded, ESCRITORIO)
            await _pestana(page, "Month")
            await _periodo(page, "2026-09-01")

            # Supervisor: borra el subgrupo y conserva el mes elegido.
            await _grupo(page, "2026-09-14")
            async with page.expect_response(
                lambda r: _es_informe(r, supervisor_user_id=admin, date="2026-09-01")
            ):
                await page.get_by_label("Supervisor").select_option(str(admin))
            await expect(_elegidas(page)).to_have_count(0)
            await expect(_alcance(page)).to_have_text("Showing whole month")
            await expect(page.locator("#report-period")).to_have_value("2026-09-01")

            # Periodo concreto: borra el subgrupo.
            await page.get_by_label("Supervisor").select_option("all")
            await expect(page.locator('[data-group-start="2026-09-14"]')).to_be_visible()
            await _grupo(page, "2026-09-14")
            await _periodo(page, "2026-01-01")
            await expect(_elegidas(page)).to_have_count(0)
            await expect(_alcance(page)).to_have_text("Showing whole month")

            # Recarga: el periodo elegido sigue, sin subgrupo.
            await page.reload()
            await expect(page.locator("#report-period")).to_have_value("2026-01-01", timeout=20_000)
            await expect(page.get_by_text("All Supervisors · January 2026")).to_be_visible()
            await expect(page.locator('[data-kpi="Total miles"]')).to_contain_text("7.0")
            await expect(_elegidas(page)).to_have_count(0)
            await contexto.close()
        finally:
            await navegador.close()


async def test_el_periodo_en_curso_lo_decide_el_servidor_no_el_navegador(
    seeded, alpha_client, live_server
):
    """Con el reloj del navegador en enero de 2024, se abre el periodo en curso
    del servidor y la primera lectura no lleva fecha."""
    from playwright.async_api import async_playwright

    await _historia(alpha_client, seeded)

    async with async_playwright() as p:
        navegador = await lanzar_edge(p)
        try:
            contexto = await navegador.new_context(base_url=live_server, viewport=ESCRITORIO)
            page = await contexto.new_page()
            await page.clock.install(time=datetime(2024, 1, 10, 12, 0))
            await abrir_sesion(page, seeded.alpha.users["route_admin"].email)
            async with page.expect_response(lambda r: _es_informe(r)) as info:
                await page.goto(RUTA)
            primera = await info.value
            assert "date=" not in primera.url
            cuerpo = await primera.json()
            assert not cuerpo["start"].startswith("2024")
            await expect(page.locator("#report-period")).to_have_value(cuerpo["start"])
            await expect(page.get_by_text("2024")).to_have_count(0)
            await contexto.close()
        finally:
            await navegador.close()


async def test_movil_selector_desglose_y_sin_desbordamiento(seeded, alpha_client, live_server):
    from playwright.async_api import async_playwright

    DESTINO.mkdir(parents=True, exist_ok=True)
    await _historia(alpha_client, seeded)

    async with async_playwright() as p:
        navegador = await lanzar_edge(p)
        try:
            contexto, page = await _abrir(navegador, live_server, seeded, MOVIL, tactil=True)
            for control in (page.get_by_label("Supervisor"), page.locator("#report-period")):
                await expect(control).to_be_visible()
            await expect(page.get_by_role("button", name="Week", exact=True)).to_be_visible()
            await _sin_desbordar(page, "inicio")

            await _periodo(page, "2026-09-14")
            antes = await _kpis(page)
            fila = page.locator('[data-group-start="2026-09-15"]')
            caja = await fila.bounding_box()
            assert caja and caja["height"] >= 44, "fila cómoda al tacto"
            async with page.expect_response(lambda r: _es_informe(r, group="2026-09-15")):
                await fila.tap()
            await expect(_alcance(page)).to_have_text("Showing Tue Sep 15 only")
            await expect(fila).to_have_attribute("aria-pressed", "true")
            await expect(page.get_by_text("Client Visit", exact=True)).to_be_visible()
            assert await _kpis(page) == antes
            await _sin_desbordar(page, "semana con día elegido")
            await page.screenshot(path=str(DESTINO / "movil-390-week-dia.png"), full_page=True)

            await _pestana(page, "Month")
            await expect(_elegidas(page)).to_have_count(0)
            await _periodo(page, "2026-09-01")
            await _grupo(page, "2026-09-14")
            await _sin_desbordar(page, "mes con semana elegida")
            await page.screenshot(path=str(DESTINO / "movil-390-month-semana.png"), full_page=True)

            await _pestana(page, "Year")
            await expect(_elegidas(page)).to_have_count(0)
            await _ano_2026(page)
            await _grupo(page, "2026-01-01")
            await _sin_desbordar(page, "año con mes elegido")
            await page.screenshot(path=str(DESTINO / "movil-390-year-mes.png"), full_page=True)
            await contexto.close()
        finally:
            await navegador.close()
