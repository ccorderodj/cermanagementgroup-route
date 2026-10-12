"""
El Excel de Reports (RTE09 §15, autorización §17).

Misma verdad que la pantalla
----------------------------
Recibe las **mismas filas** que produjeron la vista —`service.construir`—, así
que no hay un segundo motor de cálculo: la suma de *Trip Miles* es el total de
millas de la pantalla, y la de *Estimated Fuel Cost* su coste estimado.

Una fila por viaje
------------------
Es el grano donde nada se repite. Las actividades seleccionadas en una parada
van **juntas** en *Activity Type*, detrás de su propósito
(`Client Visit: Service Review, Staffing Follow-up`): una fila por actividad
multiplicaría las millas y la duración del viaje.

Lo que no se inventa
--------------------
* Kilometraje pendiente o sin cifra: el texto `Pending` / `Unresolved`, no `0`.
* Precio de referencia no disponible: **celda vacía**. Nunca un `0` para que la
  multiplicación funcione. Eso es lo que distingue una fila en espera —precio
  vacío, coste `0` visible por decisión de producto— de un cero calculado, que
  sí lleva su precio.
* Ni coordenadas, ni credenciales: no están en las filas.
"""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font

from app.routers_api.worksessions.time_zones import zona_valida

#: Las 16 columnas aprobadas, en su orden. Ni una más.
COLUMNAS: tuple[str, ...] = (
    "Date",
    "Supervisor",
    "Vehicle",
    "Activity Type",
    "Reference",
    "Start",
    "End",
    "Duration",
    "Trip Miles",
    "Operational MPG",
    "Fuel Grade",
    "Reference Fuel Price",
    "Estimated Gallons",
    "Estimated Fuel Cost",
    "Outcome",
    "Notes",
)
PROPOSITO: dict[str, str] = {
    "client_visit": "Client Visit",
    "recruiting": "Recruiting",
    "employee_visit": "Employee Visit",
    "check_delivery": "Check Delivery",
    "office": "Office",
    "other": "Other",
    "home": "Home",
}

GRADO: dict[str, str] = {
    "regular": "Regular",
    "midgrade": "Midgrade",
    "premium": "Premium",
    "diesel": "Diesel",
}

ESTADO_COMBUSTIBLE: dict[str, str] = {
    "calculated": "Calculated",
    "partial": "Partially calculated",
    "waiting_reference": "Waiting for fuel reference",
    "missing_input": "Missing input",
    "none": "No trips",
}


def reloj(instante: datetime | None, zona: str | None, desfase: int | None) -> str:
    """`5:41 PM` en la zona de la jornada: la misma regla que la pantalla.

    Zona IANA si la hay; si no, el desfase con que se registró; si tampoco, UTC
    y lo dice. Sin sufijo de zona en el caso normal (ajuste AM/PM del PO).
    """
    if instante is None:
        return ""
    zona = zona_valida(zona)
    if zona is not None:
        local = instante.astimezone(ZoneInfo(zona))
        sufijo = ""
    elif desfase is not None:
        local = instante.astimezone(timezone(timedelta(minutes=desfase)))
        sufijo = ""
    else:
        local = instante.astimezone(timezone.utc)
        sufijo = " UTC"
    hora = local.hour % 12 or 12
    return f"{hora}:{local.minute:02d} {'AM' if local.hour < 12 else 'PM'}{sufijo}"


def duracion(segundos: int | None) -> str:
    """`1h 05m` / `42m`, la forma del explorador. Vacío si no hubo bloque."""
    if segundos is None:
        return ""
    h, m = divmod(round(segundos / 60), 60)
    return f"{h}h {m:02d}m" if h else f"{m}m"


def _tipo(fila: dict) -> str:
    proposito = PROPOSITO.get(fila["purpose"], fila["purpose"])
    if fila["activity_labels"]:
        return f"{proposito}: {', '.join(fila['activity_labels'])}"
    return proposito


def _millas(fila: dict):
    if fila["official_miles"] is not None:
        return float(fila["official_miles"])
    if fila["pendiente"]:
        return "Pending"
    if fila["sin_cifra"]:
        return "Unresolved"
    return ""


def _fin(fila: dict) -> datetime | None:
    # Con parada, su cierre; sin ella —un regreso a casa—, la llegada.
    if fila["execution_id"] is not None:
        return fila["activity_ended_at"]
    return fila["arrived_at"]


def _num(valor: Decimal | None, decimales: int):
    if valor is None:
        return None
    return float(Decimal(valor).quantize(Decimal(1).scaleb(-decimales), rounding=ROUND_HALF_UP))


def libro(*, vista: dict, filas: list[dict], supervisor: str) -> bytes:
    """El `.xlsx`: el detalle y un resumen que concilia con la pantalla."""
    wb = Workbook()
    detalle = wb.active
    detalle.title = "Activity Detail"
    detalle.append(list(COLUMNAS))
    for celda in detalle[1]:
        celda.font = Font(bold=True)

    for f in filas:
        en_curso = f["execution_id"] is not None and f["activity_ended_at"] is None
        detalle.append(
            [
                f["session_date"],
                f["supervisor_name"],
                f["vehicle_label"] or "",
                _tipo(f),
                f["reference"] or "",
                reloj(f["trip_started_at"], f["time_zone"], f["utc_offset_minutes"]),
                "In progress"
                if en_curso
                else reloj(_fin(f), f["time_zone"], f["utc_offset_minutes"]),
                duracion(f["segundos"] if f["execution_id"] is not None and not en_curso else None),
                _millas(f),
                _num(f["mpg"], 2),
                GRADO.get(f["fuel_grade"] or "", ""),
                _num(f["reference_price"], 4),
                _num(f["estimated_gallons"], 3),
                _num(f["estimated_fuel_cost"], 2),
                f["outcome_label"] or "",
                f["notes"] or "",
            ]
        )
    for celda in detalle["A"][1:]:
        celda.number_format = "yyyy-mm-dd"

    resumen = wb.create_sheet("Summary")
    s = vista["summary"]
    for etiqueta, valor in (
        ("Supervisor", supervisor),
        ("Period", f"{vista['start'].isoformat()} – {vista['end'].isoformat()}"),
        ("Total miles", float(s["official_miles"])),
        ("Mileage pending", "Yes" if s["mileage_pending"] else "No"),
        ("Trips without mileage", s["mileage_unresolved"]),
        ("Activity time", duracion(s["activity_seconds"])),
        ("Activities", s["activities"]),
        ("Est. fuel cost", float(s["estimated_fuel_cost"])),
        (
            "Fuel estimate status",
            ESTADO_COMBUSTIBLE.get(s["fuel_cost_status"], s["fuel_cost_status"]),
        ),
    ):
        resumen.append([etiqueta, valor])
        resumen.cell(row=resumen.max_row, column=1).font = Font(bold=True)

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
