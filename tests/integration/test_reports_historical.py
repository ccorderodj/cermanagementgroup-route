"""
RTE09 E01: periodo histórico y desglose por subgrupo.

Qué se defiende
---------------
* Las opciones del selector de periodo salen del servidor: empiezan en el
  periodo **en curso del servidor**, retroceden hasta el primer dato del
  alcance autorizado (tenant + supervisor) y tienen tope.
* `group` acota **sólo** el desglose: KPIs, grupos y Excel siguen siendo del
  periodo principal.
* El subgrupo concilia con su fila de la tarjeta izquierda: mismas millas,
  tiempo, eventos y coste, sin duplicar por actividades.
* Un `group` que no es un día / semana / mes del periodo se rechaza.

Fechas fijas y `date` explícito: el resultado no depende del día de ejecución,
salvo las pruebas del periodo en curso, que lo comparan con lo que el propio
servidor dice que es hoy.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.routers_api.activityexplorer.dao import periodo
from app.routers_api.reports.service import OPCIONES_MAXIMAS
from tests.integration.test_activity_explorer import (
    FOTO,
    _fechar_jornada,
    _jornada_con_parada,
    _perfil_y_vehiculo,
)
from tests.integration.test_reports import (
    MIERCOLES,
    RUTA,
    _d,
    _estado,
    _informe,
    _jornada,
    _millas,
    _precio,
)

pytestmark = pytest.mark.integration


def _sin_desglose(vista: dict) -> dict:
    """La vista sin lo que el subgrupo puede cambiar."""
    return {
        k: v
        for k, v in vista.items()
        if k not in ("breakdown", "breakdown_excluded", "activity_types", "selected_group")
    }


def _fila(vista: dict, inicio: str) -> dict:
    return next(g for g in vista["groups"] if g["start"] == inicio)


def _total_desglose(vista: dict, campo: str) -> Decimal:
    return sum(
        (_d(r[campo]) for r in vista["breakdown"] + vista["breakdown_excluded"]),
        Decimal("0"),
    )


# ── Selector de periodo ─────────────────────────────────────────────────────


async def test_por_defecto_el_periodo_en_curso_lo_decide_el_servidor(seeded, alpha_client):
    """Sin `date`, el periodo en curso; y la primera opción es ese mismo."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-D")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0")

    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    for rango in ("week", "month", "year"):
        vista = (await alpha_client.get(RUTA, params={"range": rango})).json()
        assert vista["periods"][0] == {"start": vista["start"], "end": vista["end"]}
        assert vista["selected_group"] is None


async def test_las_opciones_llegan_hasta_el_primer_dato_y_no_mas(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-O")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0")

    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    vista = (await alpha_client.get(RUTA, params={"range": "month"})).json()
    ultimo = vista["periods"][-1]
    assert ultimo == {"start": "2026-09-01", "end": "2026-09-30"}
    inicios = [p["start"] for p in vista["periods"]]
    assert inicios == sorted(inicios, reverse=True), "de la más reciente a la más antigua"
    assert len(inicios) == len(set(inicios))


async def test_sin_datos_solo_el_periodo_en_curso(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-N")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    vista = (await alpha_client.get(RUTA, params={"range": "year"})).json()
    assert vista["periods"] == [{"start": vista["start"], "end": vista["end"]}]


async def test_las_opciones_tienen_tope_y_conservan_el_periodo_mostrado(seeded, alpha_client):
    """Datos de hace años: 52 semanas como mucho, y una semana antigua pedida
    explícitamente (una recarga) sigue en la lista."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-T")
    await _jornada(alpha_client, seeded, dia=date(2022, 3, 9), millas="5.0")
    # Una jornada de hoy: sin zona, el «hoy» del supervisor es su jornada más
    # reciente (T-1/T-2), así que la de 2022 sola anclaría el en curso en 2022.
    await _jornada_con_parada(alpha_client, seeded)

    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    actual = (await alpha_client.get(RUTA, params={"range": "week"})).json()
    assert len(actual["periods"]) == OPCIONES_MAXIMAS["week"]
    assert actual["periods"][0] == {"start": actual["start"], "end": actual["end"]}

    antigua = (
        await alpha_client.get(RUTA, params={"range": "week", "date": "2022-03-07"})
    ).json()
    assert {"start": "2022-03-07", "end": "2022-03-13"} in antigua["periods"]
    assert antigua["summary"]["official_miles"] == "5.0"


async def test_las_opciones_siguen_al_supervisor(seeded, alpha_client):
    """All Supervisors llega hasta el dato más antiguo de todos; uno, al suyo."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-S1")
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-S2", usuario="route_admin")
    await _jornada(alpha_client, seeded, dia=date(2026, 1, 20), millas="5.0")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="5.0", usuario="route_admin")

    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    todos = (await alpha_client.get(RUTA, params={"range": "month"})).json()
    uno = (
        await alpha_client.get(
            RUTA,
            params={"range": "month", "supervisor_user_id": seeded.alpha.users["route_admin"].id},
        )
    ).json()
    assert todos["periods"][-1]["start"] == "2026-01-01"
    assert uno["periods"][-1]["start"] == "2026-09-01"


async def test_las_opciones_no_cuentan_datos_de_otro_tenant(seeded, alpha_client, beta_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-TA")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="5.0")
    await _perfil_y_vehiculo(beta_client, seeded, unidad="V-E1-TB", compania="beta")
    await beta_client.login(seeded.beta.users["supervisor"].email)
    beta = await _jornada_minima(beta_client)
    await _fechar_jornada(beta, date(2024, 2, 7))

    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    vista = (await alpha_client.get(RUTA, params={"range": "year"})).json()
    assert vista["periods"][-1]["start"] == "2026-01-01"


async def _jornada_minima(cliente) -> int:
    """Una jornada mínima del supervisor de beta: basta con que exista."""
    jornada = (await cliente.post("/api/worksessions", json={})).json()
    assert "id" in jornada, jornada
    return jornada["id"]


async def test_un_periodo_historico_vacio_no_inventa_nada(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-V")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="5.0")
    vista = await _informe(alpha_client, seeded, range="month", date="2026-08-01")
    assert vista["groups"] == [] and vista["breakdown"] == []
    assert vista["summary"]["official_miles"] == "0.0"


# ── Subgrupo: sólo el desglose ──────────────────────────────────────────────


async def _semana(alpha_client, seeded) -> None:
    """Martes 15: cliente 10 mi con 3 actividades. Miércoles 16: oficina 3 mi."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-W")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0", actividades=3)
    await _jornada(alpha_client, seeded, dia=MIERCOLES, millas="3.0", proposito="office")


async def test_semana_dia_acota_solo_el_desglose(seeded, alpha_client):
    await _semana(alpha_client, seeded)
    entera = await _informe(alpha_client, seeded, range="week")
    dia = await _informe(alpha_client, seeded, range="week", group="2026-09-15")

    assert dia["selected_group"] == {"start": "2026-09-15", "end": "2026-09-15"}
    assert [r["purpose"] for r in dia["breakdown"]] == ["client_visit"]
    assert dia["breakdown"][0]["official_miles"] == "10.0"
    # KPIs, grupos, opciones y periodo: idénticos al periodo entero.
    assert _sin_desglose(dia) == _sin_desglose(entera)
    assert entera["summary"]["official_miles"] == "13.0"
    # Varias actividades: una parada, una vez.
    assert dia["breakdown"][0]["activities"] == 1
    assert sorted(t["count"] for t in dia["activity_types"]) == [1, 1, 1]


async def test_mes_semana_acota_solo_el_desglose(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-M")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 2), millas="7.0", proposito="office")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0")

    entera = await _informe(alpha_client, seeded, range="month")
    semana = await _informe(alpha_client, seeded, range="month", group="2026-09-01")
    assert semana["selected_group"] == {"start": "2026-09-01", "end": "2026-09-06"}
    assert [(r["purpose"], r["official_miles"]) for r in semana["breakdown"]] == [
        ("office", "7.0")
    ]
    assert _sin_desglose(semana) == _sin_desglose(entera)


async def test_ano_mes_acota_solo_el_desglose(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-Y")
    await _jornada(alpha_client, seeded, dia=date(2026, 1, 10), millas="7.0", proposito="office")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0")

    entera = await _informe(alpha_client, seeded, range="year")
    enero = await _informe(alpha_client, seeded, range="year", group="2026-01-01")
    assert enero["selected_group"] == {"start": "2026-01-01", "end": "2026-01-31"}
    assert [(r["purpose"], r["official_miles"]) for r in enero["breakdown"]] == [
        ("office", "7.0")
    ]
    assert _sin_desglose(enero) == _sin_desglose(entera)
    # Un mes sin datos del mismo año: válido, vacío.
    marzo = await _informe(alpha_client, seeded, range="year", group="2026-03-01")
    assert marzo["breakdown"] == [] and marzo["breakdown_excluded"] == []


async def test_el_subgrupo_concilia_con_su_fila(seeded, alpha_client):
    """Desglose del subgrupo (con lo excluido) = su fila de la tarjeta izquierda."""
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-C")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="24.0", actividades=2)
    await _jornada(alpha_client, seeded, dia=MIERCOLES, millas="12.0", proposito="office")
    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))

    for inicio in ("2026-09-15", "2026-09-16"):
        vista = await _informe(alpha_client, seeded, range="week", group=inicio)
        fila = _fila(vista, inicio)
        assert _total_desglose(vista, "official_miles") == _d(fila["official_miles"])
        assert _total_desglose(vista, "estimated_fuel_cost") == _d(fila["estimated_fuel_cost"])
        assert sum(r["activities"] for r in vista["breakdown"]) == fila["activities"]
        assert sum(r["activity_seconds"] for r in vista["breakdown"]) == fila["activity_seconds"]
    assert _fila(vista, "2026-09-16")["estimated_fuel_cost"] == "1.50"


async def test_un_dia_sin_actividad_solo_con_millas(seeded, alpha_client):
    """El 15, una visita; el 16, una jornada con sólo un regreso a casa: ese día
    se puede elegir, no tiene filas de visita y sus millas van aparte."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-H")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0")

    await alpha_client.login(seeded.alpha.users["supervisor"].email)
    jornada = (await alpha_client.post("/api/worksessions", json={})).json()
    await alpha_client.post(
        f"/api/odometer/sessions/{jornada['id']}/start/photo",
        files={"photo": ("odo.png", FOTO, "image/png")},
    )
    await alpha_client.post(
        f"/api/odometer/sessions/{jornada['id']}/start/confirm", json={"reading": "100200.0"}
    )
    regreso = (await alpha_client.post("/api/trips", json={"purpose": "home"})).json()
    assert "id" in regreso, regreso
    await alpha_client.post(f"/api/trips/{regreso['id']}/start", json={})
    await alpha_client.post(f"/api/trips/{regreso['id']}/arrive", json={})
    await _millas(regreso["id"], "4.0")
    await _fechar_jornada(jornada["id"], MIERCOLES)

    vista = await _informe(alpha_client, seeded, range="week", group="2026-09-16")
    fila = _fila(vista, "2026-09-16")
    assert fila["activities"] == 0 and fila["official_miles"] == "4.0"
    assert vista["breakdown"] == []
    assert [(r["purpose"], r["official_miles"]) for r in vista["breakdown_excluded"]] == [
        ("home", "4.0")
    ]
    assert vista["summary"]["official_miles"] == "14.0"


async def test_un_dia_sin_nada_es_valido_y_vacio(seeded, alpha_client):
    await _semana(alpha_client, seeded)
    vista = await _informe(alpha_client, seeded, range="week", group="2026-09-19")
    assert vista["selected_group"] == {"start": "2026-09-19", "end": "2026-09-19"}
    assert vista["breakdown"] == [] and vista["activity_types"] == []
    assert vista["summary"]["official_miles"] == "13.0"


async def test_el_subgrupo_dice_la_verdad_de_lo_pendiente(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-PE")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0")
    pendiente = await _jornada_con_parada(alpha_client, seeded)
    await _fechar_jornada(pendiente["jornada"]["id"], MIERCOLES)
    sin_cifra = await _jornada_con_parada(alpha_client, seeded)
    await _millas(sin_cifra["viaje"]["id"], "0", estado="not_calculable")
    await _fechar_jornada(sin_cifra["jornada"]["id"], MIERCOLES)

    vista = await _informe(alpha_client, seeded, range="week", group="2026-09-16")
    fila = vista["breakdown"][0]
    assert fila["official_miles"] == "0.0"
    assert fila["mileage_pending"] is True
    assert fila["mileage_unresolved"] == 1
    otro = await _informe(alpha_client, seeded, range="week", group="2026-09-15")
    assert otro["breakdown"][0]["mileage_pending"] is False


async def test_el_subgrupo_en_espera_de_precio_sigue_la_regla(seeded, alpha_client):
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-F")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="24.0")
    vista = await _informe(alpha_client, seeded, range="week", group="2026-09-15")
    fila = vista["breakdown"][0]
    assert fila["estimated_fuel_cost"] == "0.00"
    assert fila["fuel_cost_status"] == "waiting_reference"


@pytest.mark.parametrize(
    ("rango", "grupo"),
    [
        ("week", "2026-09-21"),   # el lunes siguiente: fuera de la semana
        ("month", "2026-09-16"),  # un miércoles: no empieza ninguna semana del mes
        ("month", "2026-10-05"),  # una semana de otro mes
        ("year", "2026-09-15"),   # no es el día 1 de un mes
        ("year", "2025-09-01"),   # un mes de otro año
    ],
)
async def test_un_grupo_que_no_es_del_periodo_se_rechaza(seeded, alpha_client, rango, grupo):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-X")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.get(
        RUTA, params={"range": rango, "date": MIERCOLES.isoformat(), "group": grupo}
    )
    assert r.status_code == 422, r.text


# ── Excel y seguridad ───────────────────────────────────────────────────────


async def test_el_excel_exporta_el_periodo_principal_aunque_llegue_group(seeded, alpha_client):
    await _semana(alpha_client, seeded)
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    base = {"range": "week", "date": MIERCOLES.isoformat()}
    sin = await alpha_client.get(f"{RUTA}/export", params=base)
    con = await alpha_client.get(f"{RUTA}/export", params={**base, "group": "2026-09-15"})
    assert sin.status_code == con.status_code == 200

    def filas(r):
        hoja = load_workbook(BytesIO(r.content))["Activity Detail"]
        return [list(f) for f in hoja.iter_rows(values_only=True)][1:]

    assert filas(con) == filas(sin)
    assert len(filas(con)) == 2, "los dos días de la semana, no sólo el del subgrupo"


async def test_el_excel_exporta_el_periodo_historico_elegido(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-E1-XH")
    await _jornada(alpha_client, seeded, dia=date(2026, 8, 12), millas="4.0")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="9.0")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.get(f"{RUTA}/export", params={"range": "month", "date": "2026-08-01"})
    hoja = load_workbook(BytesIO(r.content))["Activity Detail"]
    datos = list(hoja.iter_rows(values_only=True))[1:]
    assert [f[8] for f in datos] == [4.0]
    assert "2026-08-01_2026-08-31" in r.headers["content-disposition"]


async def test_group_no_abre_otro_camino_a_otro_tenant(seeded, alpha_client, beta_client):
    await _perfil_y_vehiculo(beta_client, seeded, unidad="V-E1-SB", compania="beta")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    params = {"range": "week", "date": MIERCOLES.isoformat(), "group": "2026-09-15"}
    ajeno = await alpha_client.get(
        RUTA, params={**params, "supervisor_user_id": seeded.beta.users["supervisor"].id}
    )
    falso = await alpha_client.get(RUTA, params={**params, "supervisor_user_id": 987654321})
    assert ajeno.status_code == falso.status_code == 404
    assert ajeno.json() == falso.json()

    await alpha_client.login(seeded.alpha.users["supervisor"].email)
    assert (await alpha_client.get(RUTA, params=params)).status_code == 403


def test_las_opciones_usan_los_limites_del_explorador():
    """Mismo `periodo` que Activity: la semana del 1 de octubre empieza el lunes 28."""
    assert periodo("week", date(2026, 10, 1)) == (date(2026, 9, 28), date(2026, 10, 4))
