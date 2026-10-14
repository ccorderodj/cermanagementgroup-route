"""
Jerarquía organizacional, RBAC y alcance de datos (V1.3: un rol por cargo).

El organigrama de prueba, en alpha (todo por API, como lo haría el Superadmin):

    CEO
     ├── COO
     │    ├── RM-A ── OSM-A, SUP-A
     │    └── SUP-C                (COO → Supervisor, sin RM)
     ├── RM-B ── SUP-B             (CEO → RM, sin COO)
     └── SALES-A                   (Sales, directo al CEO)

    `supervisor` (sembrado) queda **sin asignar**.

Cada persona tiene el rol de su cargo (CEO, COO, Regional Manager, OSM,
supervisor, Sales) y su superior se elige **al crearla**, en la misma
operación. Beta es el otro tenant.

Lo que se defiende: los seis roles con sus capacidades efectivas, el cargo que
sale del rol, los superiores permitidos (Sales sólo bajo un CEO), las
invariantes del modelo, el alcance por rama en Today, Activity, Reports y Excel
—COO y RM **nunca** ven a Sales—, la historia por periodo (D-3), la
administración delegada por todas las rutas de usuarios, la igualdad
OSM = Supervisor = Sales y el `route_manager` heredado.
"""

from __future__ import annotations

import asyncio
from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.database import async_session_maker
from tests.integration.conftest import TEST_PASSWORD, SeededUser, TenantClient
from tests.integration.test_activity_explorer import _perfil_y_vehiculo
from tests.integration.test_reports import _jornada

pytestmark = pytest.mark.integration

ORG = "/api/route/organization"
DIA = date(2026, 9, 16)

#: Clave de la persona → (rol, superior). En orden: cada superior antes.
ORGANIGRAMA: tuple[tuple[str, str, str | None], ...] = (
    ("ceo", "CEO", None),
    ("coo", "COO", "ceo"),
    ("rm_a", "Regional Manager", "coo"),
    ("rm_b", "Regional Manager", "ceo"),
    ("osm_a", "OSM", "rm_a"),
    ("sup_a", "supervisor", "rm_a"),
    ("sup_b", "supervisor", "rm_b"),
    ("sup_c", "supervisor", "coo"),
    ("sales_a", "Sales", "ceo"),
)
EJECUTORES = ("osm_a", "sup_a", "sup_b", "sup_c", "sales_a")


# ── Ayudas ──────────────────────────────────────────────────────────────────


async def _alta(cliente, seeded, clave, rol, superior=None, *, actor="route_admin", compania="alpha"):
    """POST de alta por el contrato de Route; devuelve la respuesta."""
    empresa = getattr(seeded, compania)
    await cliente.login(empresa.users[actor].email)
    cuerpo = {
        "username": f"{compania}_{clave}",
        "email": f"{clave}@{compania}.example.com",
        "first_name": clave.upper(),
        "last_name": "Org",
        "password": TEST_PASSWORD,
        "role_id": empresa.roles[rol],
    }
    if superior is not None:
        cuerpo["superior_user_id"] = empresa.users[superior].id
    return await cliente.post("/api/route/users", json=cuerpo)


async def _persona(cliente, seeded, clave, rol, superior=None, *, compania="alpha") -> SeededUser:
    r = await _alta(cliente, seeded, clave, rol, superior, compania=compania)
    assert r.status_code in (200, 201), (clave, r.text)
    empresa = getattr(seeded, compania)
    usuario = SeededUser(id=r.json()["id"], email=f"{clave}@{compania}.example.com",
                         username=f"{compania}_{clave}", role=rol)
    empresa.users[clave] = usuario
    return usuario


async def _organigrama(cliente, seeded, *, perfiles: bool = True) -> None:
    for clave, rol, jefe in ORGANIGRAMA:
        if rol == "CEO":
            # Un único CEO por compañía (V1.4): el del organigrama es el que
            # siembra la fixture.
            seeded.alpha.users[clave] = seeded.alpha.users["ceo_base"]
            continue
        await _persona(cliente, seeded, clave, rol, jefe)
    if perfiles:
        for clave in (*EJECUTORES, "supervisor"):
            await _perfil_y_vehiculo(cliente, seeded, unidad=f"V-ORG-{clave}", usuario=clave)


async def _superior(cliente, seeded, clave, jefe, *, actor="route_admin", **extra):
    """Cambia a quién reporta, por la edición del usuario (rol + superior)."""
    await cliente.login(seeded.alpha.users[actor].email)
    return await cliente.put(
        f"/api/route/users/{seeded.alpha.users[clave].id}",
        json={"superior_user_id": seeded.alpha.users[jefe].id if jefe else None, **extra},
    )


async def _cargo(cliente, seeded, clave, posicion, superior=None, *, actor="route_admin", **extra):
    """Por la API de organización (sigue existiendo; el cargo debe ser el del rol)."""
    await cliente.login(seeded.alpha.users[actor].email)
    return await cliente.put(
        f"{ORG}/{seeded.alpha.users[clave].id}",
        json={
            "position": posicion,
            "superior_user_id": seeded.alpha.users[superior].id if superior else None,
            **extra,
        },
    )


def _ids(seeded, *claves) -> set[int]:
    return {seeded.alpha.users[c].id for c in claves}


async def _como(cliente, seeded, clave):
    await cliente.login(seeded.alpha.users[clave].email)
    return cliente


async def _hoy(cliente, seeded, clave) -> set[int]:
    await _como(cliente, seeded, clave)
    r = await cliente.get("/api/live/today")
    assert r.status_code == 200, r.text
    return {s["user_id"] for s in r.json()["supervisors"]}


async def _explorables(cliente, seeded, clave) -> set[int]:
    await _como(cliente, seeded, clave)
    r = await cliente.get("/api/activity-explorer", params={"range": "week", "date": DIA.isoformat()})
    assert r.status_code == 200, r.text
    return {s["user_id"] for s in r.json()["supervisors"]}


async def _informe(cliente, seeded, clave, **params) -> dict:
    await _como(cliente, seeded, clave)
    params.setdefault("range", "week")
    params.setdefault("date", DIA.isoformat())
    r = await cliente.get("/api/reports", params=params)
    assert r.status_code == 200, r.text
    return r.json()


async def _filas(sql: str, **p):
    async with async_session_maker() as s:
        return (await s.execute(text(sql), p)).all()


async def _ejecutar(sql: str, **p):
    async with async_session_maker() as s:
        await s.execute(text(sql), p)
        await s.commit()


async def _vigente(seeded, clave):
    filas = await _filas(
        "SELECT position, superior_user_id, version FROM org_assignment "
        "WHERE user_id = :u AND valid_to IS NULL",
        u=seeded.alpha.users[clave].id,
    )
    return filas[0] if filas else None


async def _capacidades(role_id: int) -> set[str]:
    return {
        f.name
        for f in await _filas(
            "SELECT p.name FROM permission p JOIN role_permission rp ON rp.permission_id = p.id "
            "WHERE rp.role_id = :r AND rp.is_active",
            r=role_id,
        )
    }


async def _retrofechar(company_id: int, desde: datetime) -> None:
    await _ejecutar(
        "UPDATE org_assignment SET valid_from = :d WHERE company_id = :c AND valid_to IS NULL",
        d=desde, c=company_id,
    )


async def _iniciada(jornada_id: int, instante: datetime) -> None:
    await _ejecutar("UPDATE work_session SET started_at = :t WHERE id = :i", t=instante, i=jornada_id)


async def _trasladar_en(user_id: int, t: datetime) -> None:
    """Fija el instante del último traslado de `user_id` (preparación de prueba)."""
    await _ejecutar(
        "UPDATE org_assignment SET valid_to = :t WHERE user_id = :u AND valid_to IS NOT NULL", t=t, u=user_id
    )
    await _ejecutar(
        "UPDATE org_assignment SET valid_from = :t WHERE user_id = :u AND valid_to IS NULL", t=t, u=user_id
    )


# ── Roles y capacidades efectivas ───────────────────────────────────────────


MANDO = {
    "route.worksession.execute", "route.standardvalues.read", "route.live.read",
    "route.activity.read", "route.reports.read", "route.reports.export",
    "users.read", "users.create", "users.update",
}
EJECUCION = {"route.worksession.execute", "route.standardvalues.read"}


async def test_los_seis_roles_con_sus_capacidades_efectivas(seeded):
    for rol in ("CEO", "COO", "Regional Manager"):
        assert await _capacidades(seeded.alpha.roles[rol]) == MANDO, rol
    for rol in ("OSM", "supervisor", "Sales"):
        assert await _capacidades(seeded.alpha.roles[rol]) == EJECUCION, rol
    # El heredado conserva las de mando; `route_admin` no cambia.
    assert await _capacidades(seeded.alpha.roles["route_manager"]) == MANDO
    assert len(await _capacidades(seeded.alpha.roles["route_admin"])) == 16
    categorias = await _filas(
        "SELECT DISTINCT category FROM role WHERE company_id = :c AND name IN "
        "('CEO','COO','Regional Manager','OSM','Sales')",
        c=seeded.alpha.id,
    )
    assert {f.category for f in categorias} == {"operative"}, "management da autoridad de aprobación"


async def test_el_formulario_de_route_ofrece_los_roles_por_su_nombre(seeded, alpha_client):
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.get("/api/route/users/assignable-roles")
    assert {o["label"] for o in r.json()} == {
        "Administrador", "CEO", "COO", "Regional Manager", "OSM", "Supervisor", "Sales",
    }
    assert "route_manager" not in {o["code"] for o in r.json()}, "el heredado ya no se ofrece"


async def test_el_rol_heredado_ya_no_se_asigna(seeded, alpha_client):
    r = await _alta(alpha_client, seeded, "viejo", "route_manager")
    assert r.status_code == 403


# ── Estructura: superior por persona y por cargo ────────────────────────────


async def test_organigrama_completo_sin_niveles_ficticios_y_sin_backfill(seeded, alpha_client):
    # Sólo el CEO raíz que siembra la fixture: nada inferido.
    sembrados = {seeded.alpha.users["ceo_base"].id, seeded.beta.users["ceo_base"].id}
    assert {f.user_id for f in await _filas("SELECT user_id FROM org_assignment")} == sembrados
    await _organigrama(alpha_client, seeded, perfiles=False)
    filas = {
        f.user_id: (f.position, f.superior_user_id)
        for f in await _filas(
            "SELECT user_id, position, superior_user_id FROM org_assignment WHERE valid_to IS NULL"
        )
    }
    u = seeded.alpha.users
    assert filas[u["sales_a"].id] == ("sales", u["ceo"].id)
    assert filas[u["rm_b"].id] == ("regional_manager", u["ceo"].id)
    assert filas[u["sup_c"].id] == ("supervisor", u["coo"].id)
    assert filas[u["ceo"].id] == ("ceo", None)
    assert u["supervisor"].id not in filas


@pytest.mark.parametrize(
    ("clave", "rol", "jefe"),
    [
        ("x_sales_coo", "Sales", "coo"),          # Sales sólo bajo un CEO
        ("x_sales_rm", "Sales", "rm_a"),
        ("x_sales_sin", "Sales", None),           # y siempre con superior
        ("x_sup_osm", "supervisor", "osm_a"),     # lateral
        ("x_osm_sup", "OSM", "sup_a"),            # lateral
        ("x_rm_rm", "Regional Manager", "rm_a"),  # par
        ("x_coo_rm", "COO", "rm_a"),              # de un inferior
        ("x_coo_sin", "COO", None),               # sin superior
        ("x_sup_sales", "supervisor", "sales_a"),  # Sales no es superior de nadie
    ],
)
async def test_superiores_no_permitidos_se_rechazan_sin_crear_nada(seeded, alpha_client, clave, rol, jefe):
    await _organigrama(alpha_client, seeded, perfiles=False)
    r = await _alta(alpha_client, seeded, clave, rol, jefe)
    assert r.status_code == 422, r.text
    assert await _filas('SELECT 1 FROM "user" WHERE username = :u', u=f"alpha_{clave}") == []


async def test_superior_de_otro_tenant_o_inactivo(seeded, alpha_client, beta_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    ajeno = seeded.beta.users["ceo_base"]
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.put(
        f"/api/route/users/{seeded.alpha.users['sales_a'].id}", json={"superior_user_id": ajeno.id}
    )
    assert r.status_code == 422
    # Un CEO suspendido no es elegible: nadie nuevo se cuelga de él, y su
    # estructura sigue intacta (V1.4, FR-05).
    await alpha_client.put(f"/api/route/users/{seeded.alpha.users['ceo'].id}/access", json={"is_active": False})
    r = await _superior(alpha_client, seeded, "sup_c", "ceo")
    assert r.status_code == 422
    assert (await _vigente(seeded, "sup_c")).superior_user_id == seeded.alpha.users["coo"].id
    assert (await _vigente(seeded, "sales_a")).superior_user_id == seeded.alpha.users["ceo"].id


async def test_el_cargo_sigue_al_rol(seeded, alpha_client):
    """No hay dos controles contradictorios: el cargo es el del rol."""
    await _organigrama(alpha_client, seeded, perfiles=False)
    assert (await _cargo(alpha_client, seeded, "sup_a", "osm", "rm_a")).status_code == 422
    assert (await _vigente(seeded, "sup_a")).position == "supervisor"


async def test_el_superadmin_queda_fuera_del_arbol(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    admin2 = await _persona(alpha_client, seeded, "admin2", "route_admin")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.put(f"{ORG}/{admin2.id}", json={
        "position": "supervisor", "superior_user_id": seeded.alpha.users["rm_a"].id})
    assert r.status_code == 422
    assert await _filas("SELECT 1 FROM org_assignment WHERE user_id = :u", u=admin2.id) == []


async def test_cambiar_rol_y_superior_es_atomico(seeded, alpha_client):
    """Pasar SUP-A a Sales sin cambiar de superior (RM-A) no es válido: no cambia
    ni el rol ni el superior."""
    await _organigrama(alpha_client, seeded, perfiles=False)
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    sup_a = seeded.alpha.users["sup_a"].id
    r = await alpha_client.put(f"/api/route/users/{sup_a}", json={"role_id": seeded.alpha.roles["Sales"]})
    assert r.status_code == 422
    rol = await _filas(
        "SELECT r.name FROM user_company uc JOIN role r ON r.id = uc.role_id "
        "WHERE uc.user_id = :u AND uc.company_id = :c", u=sup_a, c=seeded.alpha.id,
    )
    assert rol[0].name == "supervisor"
    assert (await _vigente(seeded, "sup_a")).position == "supervisor"
    # Con el superior adecuado, sí: rol y cargo cambian juntos.
    r = await alpha_client.put(f"/api/route/users/{sup_a}", json={
        "role_id": seeded.alpha.roles["Sales"], "superior_user_id": seeded.alpha.users["ceo"].id})
    assert r.status_code == 200, r.text
    vigente = await _vigente(seeded, "sup_a")
    assert (vigente.position, vigente.superior_user_id) == ("sales", seeded.alpha.users["ceo"].id)
    eventos = await _filas(
        "SELECT changes FROM audit_event WHERE entity_type = 'org_assignment' "
        "AND entity_id = :u AND action = 'reassign'", u=sup_a,
    )
    assert eventos[-1].changes["position"] == {"old": "supervisor", "new": "sales"}


async def test_un_manager_con_subordinados_no_baja_ni_se_retira(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    rm_a = seeded.alpha.users["rm_a"].id
    r = await alpha_client.put(f"/api/route/users/{rm_a}", json={"role_id": seeded.alpha.roles["supervisor"]})
    assert r.status_code == 409
    assert (await alpha_client.delete(f"{ORG}/{rm_a}")).status_code == 409
    assert (await alpha_client.delete(f"/api/route/users/{rm_a}")).status_code == 409
    # Un CEO con Sales directo no puede pasar a COO.
    ceo = seeded.alpha.users["ceo"].id
    r = await alpha_client.put(f"/api/route/users/{ceo}", json={"role_id": seeded.alpha.roles["COO"]})
    assert r.status_code in (409, 422)


async def test_pasar_a_ceo_con_un_ceo_vigente_es_409(seeded, alpha_client):
    """Un único CEO por compañía (V1.4): ascender a otro a CEO no crea un
    segundo; se usa la sustitución. Nada cambia."""
    await _organigrama(alpha_client, seeded, perfiles=False)
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.put(
        f"/api/route/users/{seeded.alpha.users['sup_c'].id}", json={"role_id": seeded.alpha.roles["CEO"]}
    )
    assert r.status_code == 409, r.text
    assert "Replace manager" in r.json()["detail"]
    vigente = await _vigente(seeded, "sup_c")
    assert (vigente.position, vigente.superior_user_id) == ("supervisor", seeded.alpha.users["coo"].id)
    rol = await _filas(
        "SELECT r.name FROM user_company uc JOIN role r ON r.id = uc.role_id WHERE uc.user_id = :u",
        u=seeded.alpha.users["sup_c"].id,
    )
    assert rol[0].name == "supervisor", "el rol tampoco cambió"


async def test_nadie_se_asigna_a_si_mismo(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, "rm_a")
    r = await alpha_client.put(f"{ORG}/{seeded.alpha.users['rm_a'].id}", json={
        "position": "regional_manager", "superior_user_id": seeded.alpha.users["ceo"].id})
    assert r.status_code in (403, 404)
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.put(f"{ORG}/{seeded.alpha.users['route_admin'].id}", json={"position": "ceo"})
    assert r.status_code in (403, 422)


async def test_reasignar_cierra_y_abre_sin_solape_y_audita(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    assert (await _superior(alpha_client, seeded, "sup_a", "rm_b")).status_code == 200
    filas = await _filas(
        "SELECT superior_user_id, valid_from, valid_to, version FROM org_assignment "
        "WHERE user_id = :u ORDER BY valid_from",
        u=seeded.alpha.users["sup_a"].id,
    )
    assert len(filas) == 2
    assert filas[0].valid_to == filas[1].valid_from, "contiguas: sin hueco ni solape"
    assert filas[1].valid_to is None and filas[1].version == filas[0].version + 1
    eventos = await _filas(
        "SELECT action, actor_user_id, changes FROM audit_event WHERE entity_type = "
        "'org_assignment' AND entity_id = :u ORDER BY id",
        u=seeded.alpha.users["sup_a"].id,
    )
    assert [e.action for e in eventos] == ["assign", "reassign"]
    assert eventos[1].actor_user_id == seeded.alpha.users["route_admin"].id
    assert eventos[1].changes["superior_user_id"] == {
        "old": seeded.alpha.users["rm_a"].id, "new": seeded.alpha.users["rm_b"].id,
    }


async def test_la_base_rechaza_dos_asignaciones_abiertas(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    with pytest.raises(Exception, match="uq_org_assignment_current|ex_org_assignment_no_overlap"):
        await _ejecutar(
            "INSERT INTO org_assignment (company_id, user_id, position, valid_from) "
            "VALUES (:c, :u, 'supervisor', now())",
            c=seeded.alpha.id, u=seeded.alpha.users["sup_a"].id,
        )


async def test_dos_reasignaciones_simultaneas_una_gana(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    version = (await _vigente(seeded, "sup_a")).version

    async def intento(jefe: str) -> int:
        async with TenantClient("alpha") as c:
            r = await _cargo(c, seeded, "sup_a", "supervisor", jefe, version=version)
            return r.status_code

    codigos = sorted(await asyncio.gather(intento("rm_b"), intento("coo")))
    assert codigos == [200, 409]
    abiertas = await _filas(
        "SELECT 1 FROM org_assignment WHERE user_id = :u AND valid_to IS NULL",
        u=seeded.alpha.users["sup_a"].id,
    )
    assert len(abiertas) == 1


# ── Alcance: Today, Activity, Reports, Excel ────────────────────────────────


async def test_today_por_rama_y_sales_solo_para_el_ceo(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    assert await _hoy(alpha_client, seeded, "rm_a") == _ids(seeded, "osm_a", "sup_a")
    assert await _hoy(alpha_client, seeded, "rm_b") == _ids(seeded, "sup_b")
    assert await _hoy(alpha_client, seeded, "coo") == _ids(seeded, "osm_a", "sup_a", "sup_c")
    assert await _hoy(alpha_client, seeded, "ceo") == _ids(seeded, *EJECUTORES)
    assert await _hoy(alpha_client, seeded, "route_admin") == _ids(seeded, *EJECUTORES, "supervisor")


async def test_activity_por_rama_y_404_fuera(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    assert await _explorables(alpha_client, seeded, "coo") == _ids(seeded, "osm_a", "sup_a", "sup_c")
    assert seeded.alpha.users["sales_a"].id in await _explorables(alpha_client, seeded, "ceo")
    params = {"range": "week", "date": DIA.isoformat()}
    for actor in ("coo", "rm_a"):
        await _como(alpha_client, seeded, actor)
        ajeno = await alpha_client.get(
            "/api/activity-explorer",
            params={**params, "supervisor_user_id": seeded.alpha.users["sales_a"].id},
        )
        falso = await alpha_client.get(
            "/api/activity-explorer", params={**params, "supervisor_user_id": 987654321}
        )
        assert ajeno.status_code == falso.status_code == 404
        assert ajeno.json() == falso.json()


async def test_reports_y_excel_con_el_mismo_alcance(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    await _jornada(alpha_client, seeded, millas="10.0", usuario="sup_a")
    await _jornada(alpha_client, seeded, millas="4.0", usuario="osm_a")
    await _jornada(alpha_client, seeded, millas="7.0", usuario="sup_b")
    await _jornada(alpha_client, seeded, millas="3.0", usuario="sales_a")
    await _jornada(alpha_client, seeded, millas="2.0", usuario="supervisor")

    assert (await _informe(alpha_client, seeded, "rm_a"))["summary"]["official_miles"] == "14.0"
    assert (await _informe(alpha_client, seeded, "coo"))["summary"]["official_miles"] == "14.0"
    assert (await _informe(alpha_client, seeded, "ceo"))["summary"]["official_miles"] == "24.0"
    assert (await _informe(alpha_client, seeded, "route_admin"))["summary"]["official_miles"] == "26.0"

    params = {"range": "week", "date": DIA.isoformat()}
    for actor in ("coo", "rm_a"):
        await _como(alpha_client, seeded, actor)
        for ruta in ("/api/reports", "/api/reports/export"):
            r = await alpha_client.get(
                ruta, params={**params, "supervisor_user_id": seeded.alpha.users["sales_a"].id}
            )
            assert r.status_code == 404, (actor, ruta)
        excel = await alpha_client.get("/api/reports/export", params=params)
        filas = list(load_workbook(BytesIO(excel.content))["Activity Detail"].iter_rows(values_only=True))[1:]
        assert 3.0 not in [f[8] for f in filas], f"{actor}: el Excel no trae a Sales"
    await _como(alpha_client, seeded, "ceo")
    excel = await alpha_client.get("/api/reports/export", params=params)
    filas = list(load_workbook(BytesIO(excel.content))["Activity Detail"].iter_rows(values_only=True))[1:]
    assert sorted(f[8] for f in filas) == [3.0, 4.0, 7.0, 10.0]


async def test_un_manager_sin_cargo_no_ve_a_nadie(seeded, alpha_client):
    """Un `route_manager` heredado sin asignación: nada, nunca el tenant."""
    await _organigrama(alpha_client, seeded)
    huerfano = await _persona(alpha_client, seeded, "huerfano", "OSM", "rm_a")
    await _ejecutar(
        "UPDATE user_company SET role_id = :r WHERE user_id = :u AND company_id = :c",
        r=seeded.alpha.roles["route_manager"], u=huerfano.id, c=seeded.alpha.id,
    )
    await _ejecutar("DELETE FROM org_assignment WHERE user_id = :u", u=huerfano.id)
    await _jornada(alpha_client, seeded, millas="10.0", usuario="sup_a")
    assert await _hoy(alpha_client, seeded, "huerfano") == set()
    assert await _explorables(alpha_client, seeded, "huerfano") == set()
    vista = await _informe(alpha_client, seeded, "huerfano")
    assert vista["supervisors"] == [] and vista["summary"]["official_miles"] == "0.0"
    await _como(alpha_client, seeded, "huerfano")
    r = await alpha_client.get("/api/users/pagination", params={"page": 1, "page_size": 50})
    assert r.json()["count"] == 0


async def test_route_manager_heredado_ve_su_rama_no_el_tenant(seeded, alpha_client):
    """Quien ya tenía `route_manager` con cargo conserva el alcance de su rama."""
    await _organigrama(alpha_client, seeded)
    await _ejecutar(
        "UPDATE user_company SET role_id = :r WHERE user_id = :u AND company_id = :c",
        r=seeded.alpha.roles["route_manager"], u=seeded.alpha.users["rm_b"].id, c=seeded.alpha.id,
    )
    assert await _hoy(alpha_client, seeded, "rm_b") == _ids(seeded, "sup_b")


async def test_osm_supervisor_y_sales_son_iguales(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    permisos = {}
    for clave in ("osm_a", "sup_a", "sales_a"):
        await _como(alpha_client, seeded, clave)
        permisos[clave] = sorted(
            f.name for f in await _filas(
                "SELECT p.name FROM permission p JOIN role_permission rp ON rp.permission_id = p.id "
                "JOIN user_company uc ON uc.role_id = rp.role_id WHERE uc.user_id = :u",
                u=seeded.alpha.users[clave].id,
            )
        )
        for ruta in ("/api/live/today", "/api/activity-explorer", "/api/reports",
                     "/api/users/pagination", "/api/route/users/pagination", ORG):
            assert (await alpha_client.get(ruta)).status_code == 403, (clave, ruta)
        assert (await alpha_client.get("/api/activity-explorer/me")).status_code == 200
    assert permisos["osm_a"] == permisos["sup_a"] == permisos["sales_a"]


async def test_sales_ejecuta_y_ve_solo_lo_suyo(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    hecho = await _jornada(alpha_client, seeded, millas="3.0", usuario="sales_a")
    assert hecho["viaje"]["id"]
    await _jornada(alpha_client, seeded, millas="10.0", usuario="sup_a")
    await _como(alpha_client, seeded, "sales_a")
    r = await alpha_client.get("/api/activity-explorer/me", params={"range": "week", "date": DIA.isoformat()})
    assert r.status_code == 200
    assert r.json()["supervisor_user_id"] == seeded.alpha.users["sales_a"].id
    assert [g["official_miles"] for g in r.json()["groups"]] == ["3.0"]


async def test_un_manager_en_su_activity_personal_ve_solo_lo_suyo(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    await _jornada(alpha_client, seeded, millas="10.0", usuario="sup_a")
    await _como(alpha_client, seeded, "ceo")
    r = await alpha_client.get("/api/activity-explorer/me", params={"range": "week", "date": DIA.isoformat()})
    assert r.status_code == 200
    assert r.json()["supervisor_user_id"] == seeded.alpha.users["ceo"].id
    assert r.json()["groups"] == []


# ── Historia por periodo (D-3) ──────────────────────────────────────────────


async def test_traslado_rm_a_rm_b_por_started_at(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    await _retrofechar(seeded.alpha.id, datetime(2026, 1, 1, tzinfo=timezone.utc))
    antes = await _jornada(alpha_client, seeded, dia=date(2026, 9, 14), millas="10.0", usuario="sup_a")
    cruza = await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="3.0", usuario="sup_a")
    despues = await _jornada(alpha_client, seeded, dia=date(2026, 9, 17), millas="5.0", usuario="sup_a")

    assert (await _superior(alpha_client, seeded, "sup_a", "rm_b")).status_code == 200
    u = seeded.alpha.users["sup_a"].id
    await _trasladar_en(u, datetime(2026, 9, 16, 3, 0, tzinfo=timezone.utc))
    await _iniciada(antes["jornada"]["id"], datetime(2026, 9, 14, 13, tzinfo=timezone.utc))
    await _iniciada(cruza["jornada"]["id"], datetime(2026, 9, 16, 1, tzinfo=timezone.utc))
    await _iniciada(despues["jornada"]["id"], datetime(2026, 9, 17, 13, tzinfo=timezone.utc))
    await _ejecutar(
        "UPDATE work_session SET ended_at = :t WHERE id = :i",
        t=datetime(2026, 9, 16, 6, tzinfo=timezone.utc), i=cruza["jornada"]["id"],
    )

    a = await _informe(alpha_client, seeded, "rm_a", supervisor_user_id=u)
    b = await _informe(alpha_client, seeded, "rm_b", supervisor_user_id=u)
    assert a["summary"]["official_miles"] == "13.0", "las dos que empezaron antes de T"
    assert b["summary"]["official_miles"] == "5.0", "sólo la que empezó después"
    assert u not in await _hoy(alpha_client, seeded, "rm_a")
    assert u in await _hoy(alpha_client, seeded, "rm_b")
    sesiones = await _filas("SELECT count(*) AS n FROM work_session WHERE user_id = :u", u=u)
    assert sesiones[0].n == 3


async def test_mover_una_rama_entera(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    await _retrofechar(seeded.alpha.id, datetime(2026, 1, 1, tzinfo=timezone.utc))
    previa = await _jornada(alpha_client, seeded, dia=date(2026, 9, 14), millas="10.0", usuario="sup_a")
    posterior = await _jornada(alpha_client, seeded, dia=date(2026, 9, 17), millas="5.0", usuario="sup_a")
    assert (await _superior(alpha_client, seeded, "rm_a", "ceo")).status_code == 200
    await _trasladar_en(seeded.alpha.users["rm_a"].id, datetime(2026, 9, 16, tzinfo=timezone.utc))
    await _iniciada(previa["jornada"]["id"], datetime(2026, 9, 14, 13, tzinfo=timezone.utc))
    await _iniciada(posterior["jornada"]["id"], datetime(2026, 9, 17, 13, tzinfo=timezone.utc))
    assert (await _informe(alpha_client, seeded, "coo"))["summary"]["official_miles"] == "10.0"
    assert (await _informe(alpha_client, seeded, "ceo"))["summary"]["official_miles"] == "15.0"


async def test_sales_asignado_y_desasignado_sin_fuga_historica(seeded, alpha_client):
    """SUP-C (bajo el COO) pasa a Sales bajo el CEO en T: lo anterior queda en la
    rama del COO; lo posterior, sólo para el CEO."""
    await _organigrama(alpha_client, seeded)
    await _retrofechar(seeded.alpha.id, datetime(2026, 1, 1, tzinfo=timezone.utc))
    previa = await _jornada(alpha_client, seeded, dia=date(2026, 9, 14), millas="6.0", usuario="sup_c")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    sup_c = seeded.alpha.users["sup_c"].id
    r = await alpha_client.put(f"/api/route/users/{sup_c}", json={
        "role_id": seeded.alpha.roles["Sales"], "superior_user_id": seeded.alpha.users["ceo"].id})
    assert r.status_code == 200, r.text
    await _trasladar_en(sup_c, datetime(2026, 9, 16, tzinfo=timezone.utc))
    posterior = await _jornada(alpha_client, seeded, dia=date(2026, 9, 17), millas="2.0", usuario="sup_c")
    await _iniciada(previa["jornada"]["id"], datetime(2026, 9, 14, 13, tzinfo=timezone.utc))
    await _iniciada(posterior["jornada"]["id"], datetime(2026, 9, 17, 13, tzinfo=timezone.utc))

    coo = await _informe(alpha_client, seeded, "coo", supervisor_user_id=sup_c)
    assert coo["summary"]["official_miles"] == "6.0"
    assert sup_c not in await _hoy(alpha_client, seeded, "coo")
    ceo = await _informe(alpha_client, seeded, "ceo", supervisor_user_id=sup_c)
    assert ceo["summary"]["official_miles"] == "8.0"


async def test_suspendido_sale_de_today_y_su_historia_queda(seeded, alpha_client):
    await _organigrama(alpha_client, seeded)
    await _jornada(alpha_client, seeded, millas="6.0", usuario="sup_a")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    r = await alpha_client.put(
        f"/api/users/{seeded.alpha.users['sup_a'].id}/access", json={"is_active": False}
    )
    assert r.status_code == 200
    assert seeded.alpha.users["sup_a"].id not in await _hoy(alpha_client, seeded, "route_admin")
    assert (await _informe(alpha_client, seeded, "route_admin"))["summary"]["official_miles"] == "6.0"


# ── Administración delegada: todas las rutas ────────────────────────────────


@pytest.mark.parametrize("base", ["/api/users", "/api/route/users"])
@pytest.mark.parametrize(
    ("actor", "visibles"),
    [
        ("ceo", ("coo", "rm_a", "rm_b", "osm_a", "sup_a", "sup_b", "sup_c", "sales_a")),
        ("coo", ("rm_a", "osm_a", "sup_a", "sup_c")),
        ("rm_a", ("osm_a", "sup_a")),
    ],
)
async def test_cada_manager_lista_solo_su_rama(seeded, alpha_client, base, actor, visibles):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, actor)
    r = await alpha_client.get(f"{base}/pagination", params={"page": 1, "page_size": 100})
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert {u["id"] for u in cuerpo["results"]} == _ids(seeded, *visibles)
    assert cuerpo["count"] == len(visibles), "el total tampoco cuenta a nadie más"


@pytest.mark.parametrize("base", ["/api/users", "/api/route/users"])
@pytest.mark.parametrize(
    ("actor", "objetivo"),
    [
        ("coo", "sales_a"), ("rm_a", "sales_a"), ("rm_a", "sup_b"), ("rm_a", "coo"),
        ("rm_a", "rm_b"), ("rm_a", "rm_a"), ("coo", "ceo"), ("coo", "route_admin"),
        ("ceo", "owner"),
    ],
)
async def test_el_manager_no_toca_fuera_de_su_rama(seeded, alpha_client, base, actor, objetivo):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, actor)
    uid = seeded.alpha.users[objetivo].id
    assert (await alpha_client.put(f"{base}/{uid}", json={"first_name": "X"})).status_code == 404
    assert (await alpha_client.put(f"{base}/{uid}/access", json={"is_active": False})).status_code == 404
    assert (await alpha_client.put(f"{ORG}/{uid}", json={"position": "supervisor",
            "superior_user_id": seeded.alpha.users[actor].id})).status_code == 404


async def test_el_manager_no_toca_otro_tenant(seeded, alpha_client, beta_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, "ceo")
    beta = seeded.beta.users["supervisor"].id
    for base in ("/api/users", "/api/route/users"):
        assert (await alpha_client.put(f"{base}/{beta}", json={"first_name": "X"})).status_code == 404
    assert (await alpha_client.put(f"{ORG}/{beta}", json={"position": "supervisor"})).status_code == 404


async def test_el_manager_gestiona_su_rama(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, "rm_a")
    sup_a = seeded.alpha.users["sup_a"].id
    assert (await alpha_client.put(f"/api/route/users/{sup_a}", json={"first_name": "Ana"})).status_code == 200
    assert (await alpha_client.put(f"/api/route/users/{sup_a}/access", json={"is_active": False})).status_code == 200
    assert (await alpha_client.put(f"/api/route/users/{sup_a}/access", json={"is_active": True})).status_code == 200
    # SUP-A pasa a OSM bajo el propio RM: rol y cargo juntos.
    r = await alpha_client.put(f"/api/route/users/{sup_a}", json={"role_id": seeded.alpha.roles["OSM"]})
    assert r.status_code == 200, r.text
    assert (await _vigente(seeded, "sup_a")).position == "osm"


async def test_el_manager_no_borra_ni_readmite(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, "rm_a")
    sup_a = seeded.alpha.users["sup_a"].id
    assert (await alpha_client.delete(f"/api/route/users/{sup_a}")).status_code == 403
    sup_c = seeded.alpha.users["sup_c"].id
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    assert (await alpha_client.delete(f"/api/route/users/{sup_c}")).status_code == 204
    await _como(alpha_client, seeded, "coo")
    r = await alpha_client.post(f"/api/route/users/{sup_c}/reenrollment",
                                json={"role_id": seeded.alpha.roles["supervisor"]})
    assert r.status_code == 404


@pytest.mark.parametrize(
    ("actor", "rol", "jefe", "codigo"),
    [
        ("rm_a", "supervisor", "rm_a", 200),
        ("rm_a", "OSM", "rm_a", 200),
        ("rm_a", "supervisor", "osm_a", 422),              # OSM no es superior de nadie
        ("rm_a", "Regional Manager", "rm_a", 403),         # su nivel
        ("rm_a", "supervisor", "rm_b", 403),               # rama ajena
        ("rm_a", "route_admin", "rm_a", 403),              # escalar
        ("rm_a", "Sales", "rm_a", 403),                    # Sales: sólo un CEO
        ("rm_a", "supervisor", None, 403),                 # sin superior: fuera de su rama
        ("coo", "Regional Manager", "coo", 200),
        ("coo", "Sales", "coo", 403),
        ("coo", "Sales", "ceo", 403),                      # el CEO no es de su rama
        ("coo", "COO", "coo", 403),
        ("ceo", "COO", "ceo", 200),
        ("ceo", "Sales", "ceo", 200),
        ("ceo", "Sales", "coo", 422),                      # Sales sólo bajo CEO directo
        ("ceo", "CEO", "ceo", 403),
    ],
)
async def test_alta_delegada(seeded, alpha_client, actor, rol, jefe, codigo):
    await _organigrama(alpha_client, seeded, perfiles=False)
    r = await _alta(alpha_client, seeded, f"nuevo_{actor}", rol, jefe, actor=actor)
    if codigo == 200:
        assert r.status_code in (200, 201), r.text
        fila = await _filas(
            "SELECT superior_user_id FROM org_assignment WHERE user_id = :u AND valid_to IS NULL",
            u=r.json()["id"],
        )
        assert fila[0].superior_user_id == seeded.alpha.users[jefe].id
    else:
        assert r.status_code == codigo, r.text
        assert await _filas(
            'SELECT 1 FROM "user" WHERE username = :u', u=f"alpha_nuevo_{actor}"
        ) == [], "nada se creó"


@pytest.mark.parametrize("rol", ["route_admin", "route_manager", "owner", "Sales", "COO"])
async def test_el_manager_no_escala_un_rol(seeded, alpha_client, rol):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, "rm_a")
    r = await alpha_client.put(
        f"/api/users/{seeded.alpha.users['sup_a'].id}", json={"role_id": seeded.alpha.roles[rol]}
    )
    assert r.status_code == 403
    assert (await _vigente(seeded, "sup_a")).position == "supervisor"


async def test_roles_asignables_por_nivel(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    esperado = {
        "rm_a": {"OSM", "supervisor"},
        "coo": {"Regional Manager", "OSM", "supervisor"},
        "ceo": {"COO", "Regional Manager", "OSM", "supervisor", "Sales"},
    }
    for actor, roles in esperado.items():
        await _como(alpha_client, seeded, actor)
        r = await alpha_client.get("/api/route/users/assignable-roles")
        assert {o["code"] for o in r.json()} == roles, actor


async def test_ejecutores_no_administran(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    for clave in ("osm_a", "sup_a", "sales_a"):
        await _como(alpha_client, seeded, clave)
        assert (await alpha_client.get("/api/route/users/pagination")).status_code == 403
        assert (await alpha_client.put(
            f"{ORG}/{seeded.alpha.users['sup_b'].id}", json={"position": "supervisor"}
        )).status_code == 403
        r = await alpha_client.post("/api/route/users", json={
            "username": f"alpha_x_{clave}", "email": f"x_{clave}@alpha.example.com",
            "first_name": "X", "last_name": "Y", "password": TEST_PASSWORD,
            "role_id": seeded.alpha.roles["supervisor"],
        })
        assert r.status_code == 403


async def test_el_manager_no_alcanza_candidatos_de_supervisor(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, "ceo")
    assert (await alpha_client.get("/api/supervisors/candidates")).status_code == 403


async def test_las_opciones_de_superior_respetan_rol_y_rama(seeded, alpha_client):
    await _organigrama(alpha_client, seeded, perfiles=False)
    await _como(alpha_client, seeded, "coo")
    o = (await alpha_client.get(f"{ORG}/options")).json()
    assert {s["user_id"] for s in o["superiors"]} == _ids(seeded, "coo", "rm_a", "osm_a", "sup_a", "sup_c")
    assert o["allowed_parents"]["sales"] == ["ceo"]
    assert o["role_positions"]["Sales"] == "sales"
