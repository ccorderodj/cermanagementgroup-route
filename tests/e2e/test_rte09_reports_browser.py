"""RTE09: Reports en el navegador, escritorio y móvil.

Lo que se comprueba
-------------------
* La cadena completa en escritorio: estado operativo fijado por el
  administrador → jornada → kilometraje → referencia de precio → KPI de coste,
  tarjetas *By day* y *Activity breakdown*, y el Excel que se descarga.
* En móvil (390 px), sin referencia de precio: el coste se ve `$0.00` —la
  decisión de producto—, el selector de supervisor sigue visible, las pestañas
  Week / Month / Year cambian el rango y la página no desborda.
* Elegir un supervisor pide ese `supervisor_user_id` y la leyenda lo nombra.

Reports enseña el periodo en curso, sin navegación: las jornadas son de hoy y
la referencia de precio tiene una fecha efectiva lejana para que aplique
cualquiera que sea el día de la ejecución.

Capturas en `var/screenshots/rte09/`, fuera del árbol versionado.
"""

from __future__ import annotations

import pathlib
import re
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from playwright.async_api import expect

from tests.e2e.conftest import abrir_sesion, lanzar_edge
from tests.integration.test_activity_explorer import _jornada_con_parada, _perfil_y_vehiculo
from tests.integration.test_reports import _estado, _millas, _precio

pytestmark = [pytest.mark.integration, pytest.mark.browser]

pytest.importorskip("playwright", reason="playwright no está instalado")

DESTINO = pathlib.Path("var/screenshots/rte09")
RUTA = "/admin/route/reports"
ESCRITORIO = {"width": 1280, "height": 900}
MOVIL = {"width": 390, "height": 844}


async def _abrir(navegador, live_server, seeded, viewport):
    contexto = await navegador.new_context(
        base_url=live_server, viewport=viewport, accept_downloads=True
    )
    page = await contexto.new_page()
    await abrir_sesion(page, seeded.alpha.users["route_admin"].email)
    await page.goto(RUTA)
    await expect(page.get_by_role("heading", name="Reports", exact=True)).to_have_count(
        1, timeout=20_000
    )
    return contexto, page


def _kpi(page, nombre: str):
    return page.locator(f'[data-kpi="{nombre}"]')


async def test_escritorio_la_cadena_completa_y_el_excel(seeded, alpha_client, live_server):
    from playwright.async_api import async_playwright

    DESTINO.mkdir(parents=True, exist_ok=True)
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-E2E")
    await _estado(alpha_client, seeded, perfil, "GA")
    hecho = await _jornada_con_parada(alpha_client, seeded, actividades=2)
    await _millas(hecho["viaje"]["id"], "24.0")  # MPG 24 → 1 galón
    await _precio(seeded.alpha.id, "GA", "regular", "3.2500", date(2000, 1, 1))

    async with async_playwright() as p:
        navegador = await lanzar_edge(p)
        try:
            contexto, page = await _abrir(navegador, live_server, seeded, ESCRITORIO)
            await expect(_kpi(page, "Total miles")).to_contain_text("24", timeout=20_000)
            await expect(_kpi(page, "Activities")).to_contain_text("1")
            await expect(_kpi(page, "Est. fuel cost")).to_contain_text("$3.25")
            await expect(page.get_by_text("By day", exact=True)).to_be_visible()
            await expect(page.get_by_text("Activity breakdown", exact=True)).to_be_visible()
            await expect(page.get_by_text("Client Visit").first).to_be_visible()
            await expect(page.get_by_label("Supervisor")).to_have_value("all")
            await expect(page.get_by_text(re.compile(r"^All Supervisors · "))).to_be_visible()
            await page.screenshot(path=str(DESTINO / "escritorio-1280-week.png"), full_page=True)

            async with page.expect_download() as info:
                await page.get_by_role("button", name="Export Excel").click()
            descarga = await info.value
            assert descarga.suggested_filename.startswith("CER_Route_Report_")
            contenido = pathlib.Path(await descarga.path()).read_bytes()
            hoja = load_workbook(BytesIO(contenido))["Activity Detail"]
            filas = list(hoja.iter_rows(values_only=True))
            assert len(filas[0]) == 16
            assert len(filas) == 2, "una fila por viaje, aunque haya dos actividades"
            assert filas[1][11] == 3.25 and filas[1][13] == 3.25

            # Elegir al supervisor pide su id y la leyenda lo nombra.
            usuario = seeded.alpha.users["supervisor"]
            async with page.expect_response(
                lambda r: "/api/reports" in r.url
                and f"supervisor_user_id={usuario.id}" in r.url
            ) as respuesta:
                await page.get_by_label("Supervisor").select_option(str(usuario.id))
            cuerpo = await (await respuesta.value).json()
            assert cuerpo["supervisor_user_id"] == usuario.id
            nombre = next(s["name"] for s in cuerpo["supervisors"] if s["user_id"] == usuario.id)
            await expect(page.get_by_text(re.compile(rf"^{re.escape(nombre)} · "))).to_be_visible()
            await expect(_kpi(page, "Est. fuel cost")).to_contain_text("$3.25")
            await contexto.close()
        finally:
            await navegador.close()


async def test_movil_sin_referencia_el_coste_es_cero_y_no_desborda(
    seeded, alpha_client, live_server
):
    from playwright.async_api import async_playwright

    DESTINO.mkdir(parents=True, exist_ok=True)
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-MOV")
    await _estado(alpha_client, seeded, perfil, "GA")
    hecho = await _jornada_con_parada(alpha_client, seeded)
    await _millas(hecho["viaje"]["id"], "24.0")
    # Un precio de otro estado: no aplica, no hay alternativa regional.
    await _precio(seeded.alpha.id, "TX", "regular", "9.0000", date(2000, 1, 1))

    async with async_playwright() as p:
        navegador = await lanzar_edge(p)
        try:
            contexto, page = await _abrir(navegador, live_server, seeded, MOVIL)
            await expect(_kpi(page, "Total miles")).to_contain_text("24", timeout=20_000)
            await expect(_kpi(page, "Est. fuel cost")).to_contain_text("$0.00")
            await expect(page.get_by_label("Supervisor")).to_be_visible()
            await page.screenshot(path=str(DESTINO / "movil-390-week-esperando.png"), full_page=True)

            for rango, agrupado in (("Month", "By week"), ("Year", "By month"), ("Week", "By day")):
                async with page.expect_response(
                    lambda r, rg=rango.lower(): "/api/reports" in r.url and f"range={rg}" in r.url
                ) as respuesta:
                    await page.get_by_role("button", name=rango, exact=True).click()
                cuerpo = await (await respuesta.value).json()
                assert cuerpo["summary"]["fuel_cost_status"] == "waiting_reference"
                await expect(page.get_by_text(agrupado, exact=True)).to_be_visible()
                await expect(_kpi(page, "Est. fuel cost")).to_contain_text("$0.00")
                ancho = await page.evaluate(
                    "document.documentElement.scrollWidth - document.documentElement.clientWidth"
                )
                assert ancho <= 0, f"{rango}: desborda {ancho}px"
            await page.screenshot(path=str(DESTINO / "movil-390-week-tras-pestanas.png"), full_page=True)
            await contexto.close()
        finally:
            await navegador.close()


async def test_supervisor_no_ve_reports(seeded, alpha_client, live_server):
    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        navegador = await lanzar_edge(p)
        try:
            contexto = await navegador.new_context(base_url=live_server, viewport=ESCRITORIO)
            page = await contexto.new_page()
            await abrir_sesion(page, seeded.alpha.users["supervisor"].email)
            respuesta = await page.goto(RUTA)
            assert respuesta.status == 403
            await contexto.close()
        finally:
            await navegador.close()
