"""
RTE09: Reports consolidados y estimación de combustible.

Qué se defiende
---------------
* Un solo modelo de lectura: KPIs, grupos, desglose y Excel salen de las mismas
  filas, y concilian entre sí y con el explorador.
* Millas: la regla compartida, cada viaje una vez —ni los tramos ni las
  actividades seleccionadas la multiplican—, y lo pendiente se dice aparte.
* Combustible: `millas / MPG congelado × precio de la referencia aplicable`,
  calculado **por jornada**; la referencia se elige por estado + grado + fecha
  efectiva; sin referencia, coste visible `0` con estado `waiting_reference`,
  que no se congela y se resuelve cuando el precio llega; una estimación final
  no cambia con precios posteriores.
* Snapshots: estado operativo y grado congelados al empezar la jornada.
* Permisos y tenant.

Las pruebas usan fechas fijas y el parámetro `date` del API: el resultado no
depende del día en que se ejecuten.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.database import async_session_maker
from tests.integration.test_activity_explorer import (
    FOTO,
    _explorar,
    _fechar_jornada,
    _jornada_con_parada,
    _perfil_y_vehiculo,
    _valores,
)

pytestmark = pytest.mark.integration

RUTA = "/api/reports"
#: Una milla oficial exacta, en metros.
MILLA = Decimal("1609.344")
#: Miércoles: la semana es del lunes 14 al domingo 20 de septiembre de 2026.
MIERCOLES = date(2026, 9, 16)


# ── Ayudas ──────────────────────────────────────────────────────────────────


async def _millas(trip_id: int, millas: str, estado: str = "calculated") -> None:
    """Fija el kilometraje del viaje, como lo dejaría el motor."""
    async with async_session_maker() as s:
        if estado == "calculated":
            await s.execute(
                text(
                    "UPDATE trip_mileage SET state = 'calculated', total_meters = :m, "
                    "calculated_at = now() WHERE trip_id = :t"
                ),
                {"m": Decimal(millas) * MILLA, "t": trip_id},
            )
        elif estado == "not_calculable":
            await s.execute(
                text(
                    "UPDATE trip_mileage SET state = 'not_calculable', "
                    "terminal_reason = 'arrival_waypoint_missing' WHERE trip_id = :t"
                ),
                {"t": trip_id},
            )
        await s.commit()


async def _precio(
    company_id: int, estado: str, grado: str, precio: str, efectiva: date,
    fuente: str = "Test feed",
) -> int:
    """Una fila de referencia, como la escribiría el trabajo programado."""
    async with async_session_maker() as s:
        nuevo = await s.scalar(
            text(
                "INSERT INTO fuel_reference(company_id, state_code, fuel_grade, "
                "price_per_gallon, effective_date, source) VALUES "
                "(:c, :e, :g, :p, :f, :s) RETURNING id"
            ),
            {"c": company_id, "e": estado, "g": grado, "p": precio, "f": efectiva, "s": fuente},
        )
        await s.commit()
        return nuevo


async def _estado(cliente, seeded, perfil: dict, codigo: str | None) -> None:
    await cliente.login(seeded.alpha.users["route_admin"].email)
    r = await cliente.put(
        f"/api/supervisors/{perfil['id']}/operational-state",
        json={"operational_state": codigo},
    )
    assert r.status_code == 200, r.text


async def _vehiculo(cliente, seeded, vehiculo: dict, **cambios) -> None:
    await cliente.login(seeded.alpha.users["route_admin"].email)
    r = await cliente.put(f"/api/vehicles/{vehiculo['id']}", json=cambios)
    assert r.status_code == 200, r.text


async def _informe(cliente, seeded, **parametros) -> dict:
    await cliente.login(seeded.alpha.users["route_admin"].email)
    parametros.setdefault("date", MIERCOLES.isoformat())
    r = await cliente.get(RUTA, params=parametros)
    assert r.status_code == 200, r.text
    return r.json()


async def _jornada(
    cliente, seeded, *, dia: date = MIERCOLES, millas: str = "10.0",
    proposito: str = "client_visit", actividades: int = 1, cerrar: bool = True,
    usuario: str = "supervisor",
) -> dict:
    """Una jornada con una parada, su kilometraje y su día de negocio."""
    hecho = await _jornada_con_parada(
        cliente, seeded, proposito=proposito, actividades=actividades,
        cerrar_jornada=cerrar, usuario=usuario,
    )
    await _millas(hecho["viaje"]["id"], millas)
    await _fechar_jornada(hecho["jornada"]["id"], dia)
    return hecho


def _d(valor) -> Decimal:
    return Decimal(str(valor))


# ── Reports: periodos, filtros y conciliación ───────────────────────────────


async def test_semana_mes_y_ano_agrupan_como_el_explorador(seeded, alpha_client):
    """Week → días, Month → semanas recortadas, Year → meses; y Σ grupos = total."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-P")
    for dia, millas in (
        (date(2026, 9, 15), "10.0"),
        (date(2026, 9, 16), "20.0"),
        (date(2026, 9, 2), "5.0"),
        (date(2026, 1, 10), "7.0"),
    ):
        await _jornada(alpha_client, seeded, dia=dia, millas=millas)

    semana = await _informe(alpha_client, seeded, range="week")
    assert (semana["start"], semana["end"]) == ("2026-09-14", "2026-09-20")
    assert semana["grouped_by"] == "day"
    assert [g["start"] for g in semana["groups"]] == ["2026-09-15", "2026-09-16"]
    assert semana["summary"]["official_miles"] == "30.0"

    mes = await _informe(alpha_client, seeded, range="month")
    assert mes["grouped_by"] == "week"
    assert [(g["start"], g["end"]) for g in mes["groups"]] == [
        ("2026-09-01", "2026-09-06"), ("2026-09-14", "2026-09-20"),
    ]
    assert mes["summary"]["official_miles"] == "35.0"

    ano = await _informe(alpha_client, seeded, range="year")
    assert ano["grouped_by"] == "month"
    assert [g["start"] for g in ano["groups"]] == ["2026-01-01", "2026-09-01"]

    for vista in (semana, mes, ano):
        s = vista["summary"]
        assert sum(_d(g["official_miles"]) for g in vista["groups"]) == _d(s["official_miles"])
        assert sum(g["activities"] for g in vista["groups"]) == s["activities"]
        assert sum(g["activity_seconds"] for g in vista["groups"]) == s["activity_seconds"]
        assert sum(_d(g["estimated_fuel_cost"]) for g in vista["groups"]) == _d(
            s["estimated_fuel_cost"]
        )


async def test_all_supervisors_suma_y_uno_filtra(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-A1")
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-A2", usuario="route_admin")
    await _jornada(alpha_client, seeded, millas="10.0")
    await _jornada(alpha_client, seeded, millas="4.0", usuario="route_admin")

    todos = await _informe(alpha_client, seeded, range="week")
    assert todos["supervisor_user_id"] is None
    assert todos["summary"]["official_miles"] == "14.0"
    assert todos["summary"]["activities"] == 2

    uno = await _informe(
        alpha_client, seeded, range="week",
        supervisor_user_id=seeded.alpha.users["route_admin"].id,
    )
    assert uno["supervisor_user_id"] == seeded.alpha.users["route_admin"].id
    assert uno["summary"]["official_miles"] == "4.0"
    assert uno["summary"]["activities"] == 1


async def test_un_periodo_vacio_no_inventa_nada(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-E")
    vista = await _informe(alpha_client, seeded, range="week", date="2025-03-05")
    assert vista["groups"] == [] and vista["breakdown"] == []
    s = vista["summary"]
    assert (s["official_miles"], s["activities"], s["estimated_fuel_cost"]) == ("0.0", 0, "0.00")
    assert s["fuel_cost_status"] == "none"


async def test_manda_session_date_aunque_la_jornada_cruce_medianoche(seeded, alpha_client):
    """Una jornada del domingo 20 que termina el lunes 21 es de la semana del 14."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-M")
    hecho = await _jornada(alpha_client, seeded, dia=date(2026, 9, 20), millas="6.0")
    async with async_session_maker() as s:
        await s.execute(
            text("UPDATE trip SET started_at = :t WHERE id = :i"),
            {"t": datetime(2026, 9, 21, 4, 30, tzinfo=timezone.utc), "i": hecho["viaje"]["id"]},
        )
        await s.commit()

    semana = await _informe(alpha_client, seeded, range="week")
    assert semana["summary"]["official_miles"] == "6.0"
    siguiente = await _informe(alpha_client, seeded, range="week", date="2026-09-21")
    assert siguiente["summary"]["official_miles"] == "0.0"


async def test_concilia_con_el_explorador_dia_a_dia(seeded, alpha_client):
    """Mismas filas, misma regla: las millas, paradas y tiempo de cada día."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-X")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="12.3")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 16), millas="4.4", actividades=2)

    informe = await _informe(
        alpha_client, seeded, range="week",
        supervisor_user_id=seeded.alpha.users["supervisor"].id,
    )
    explorador = await _explorar(
        alpha_client, range="week", date=MIERCOLES.isoformat(),
        supervisor_user_id=seeded.alpha.users["supervisor"].id,
    )
    for ours, theirs in zip(informe["groups"], explorador["groups"], strict=True):
        assert ours["start"] == theirs["start"]
        assert ours["official_miles"] == theirs["official_miles"]
        assert ours["activities"] == theirs["activities"]
        assert ours["activity_seconds"] == theirs["activity_seconds"]


# ── Millas ──────────────────────────────────────────────────────────────────


async def test_los_tramos_no_se_suman_al_total(seeded, alpha_client):
    """Un tramo es la procedencia del total, no millas adicionales."""
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-T")
    hecho = await _jornada(alpha_client, seeded, millas="10.0")
    async with async_session_maker() as s:
        km = await s.scalar(
            text("SELECT id FROM trip_mileage WHERE trip_id = :t"), {"t": hecho["viaje"]["id"]}
        )
        await s.execute(
            text(
                "INSERT INTO trip_mileage_segment(company_id, trip_mileage_id, sequence, "
                "from_event_kind, from_latitude, from_longitude, from_evidence_level, "
                "from_captured_at, to_event_kind, to_latitude, to_longitude, "
                "to_evidence_level, to_captured_at, distance_meters, provider, method, "
                "computed_at, created_at, updated_at) VALUES (:c, :k, 1, 'start_trip', 33.7, -84.4, "
                "'fresh', now(), 'arrived', 33.8, -84.3, 'fresh', now(), :m, 'osrm', "
                "'route', now(), now(), now())"
            ),
            {"c": seeded.alpha.id, "k": km, "m": 10 * MILLA},
        )
        await s.commit()

    vista = await _informe(alpha_client, seeded, range="week")
    assert vista["summary"]["official_miles"] == "10.0"


async def test_varias_actividades_no_multiplican_millas_ni_tiempo(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-MA")
    await _jornada(alpha_client, seeded, millas="10.0", actividades=3)

    vista = await _informe(alpha_client, seeded, range="week")
    s = vista["summary"]
    assert s["official_miles"] == "10.0"
    assert s["activities"] == 1
    explorador = await _explorar(
        alpha_client, range="week", date=MIERCOLES.isoformat(),
        supervisor_user_id=seeded.alpha.users["supervisor"].id,
    )
    assert s["activity_seconds"] == explorador["groups"][0]["activity_seconds"]
    # Cada actividad seleccionada, una ocurrencia.
    assert len(vista["activity_types"]) == 3
    assert all(t["count"] == 1 for t in vista["activity_types"])


async def test_calculado_pendiente_y_sin_cifra_se_distinguen(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-PE")
    await _jornada(alpha_client, seeded, millas="10.0")
    pendiente = await _jornada_con_parada(alpha_client, seeded)
    await _fechar_jornada(pendiente["jornada"]["id"], MIERCOLES)
    sin_cifra = await _jornada_con_parada(alpha_client, seeded)
    await _millas(sin_cifra["viaje"]["id"], "0", estado="not_calculable")
    await _fechar_jornada(sin_cifra["jornada"]["id"], MIERCOLES)

    s = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert s["official_miles"] == "10.0"
    assert s["mileage_pending"] is True
    assert s["mileage_unresolved"] == 1


# ── Actividad: desglose por propósito ───────────────────────────────────────


async def test_el_desglose_es_por_proposito_y_concilia_con_el_total(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-D")
    await _jornada(alpha_client, seeded, millas="10.0", proposito="client_visit")
    await _jornada(alpha_client, seeded, millas="3.0", proposito="office")
    casa = await _jornada_con_parada(alpha_client, seeded, cerrar_jornada=False)
    await _millas(casa["viaje"]["id"], "1.0")
    regreso = (await alpha_client.post("/api/trips", json={"purpose": "home"})).json()
    await alpha_client.post(f"/api/trips/{regreso['id']}/start", json={})
    await alpha_client.post(f"/api/trips/{regreso['id']}/arrive", json={})
    await _millas(regreso["id"], "5.0")
    await _fechar_jornada(casa["jornada"]["id"], MIERCOLES)

    vista = await _informe(alpha_client, seeded, range="week")
    filas = {r["purpose"]: r for r in vista["breakdown"]}
    assert list(filas) == ["client_visit", "office"], "en el orden del V0.7"
    assert filas["client_visit"]["official_miles"] == "11.0"
    assert filas["client_visit"]["activities"] == 2
    assert filas["office"]["official_miles"] == "3.0"
    # El regreso a casa no es una visita: no tiene fila, pero sus millas cuentan.
    assert [r["purpose"] for r in vista["breakdown_excluded"]] == ["home"]
    total = sum(_d(r["official_miles"]) for r in vista["breakdown"] + vista["breakdown_excluded"])
    assert total == _d(vista["summary"]["official_miles"]) == Decimal("19.0")


# ── Estado operativo y grado: snapshots ─────────────────────────────────────


async def test_estado_y_grado_se_congelan_al_empezar(seeded, alpha_client):
    """Editar luego el supervisor o el vehículo no cambia la jornada ni su coste."""
    perfil, vehiculo = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-S")
    await _estado(alpha_client, seeded, perfil, "GA")
    hecho = await _jornada(alpha_client, seeded, millas="24.0")  # MPG 24 → 1 galón

    await _estado(alpha_client, seeded, perfil, "TX")
    await _vehiculo(alpha_client, seeded, vehiculo, fuel_grade="diesel", operational_mpg="12.00")

    async with async_session_maker() as s:
        fila = (
            await s.execute(
                text(
                    "SELECT operational_state_snapshot, fuel_grade_snapshot, mpg_snapshot "
                    "FROM work_session WHERE id = :i"
                ),
                {"i": hecho["jornada"]["id"]},
            )
        ).one()
    assert (fila.operational_state_snapshot, fila.fuel_grade_snapshot) == ("GA", "regular")
    assert fila.mpg_snapshot == Decimal("24.00")

    # Precios para los valores de antes y de ahora: se usa el congelado.
    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))
    await _precio(seeded.alpha.id, "TX", "diesel", "9.0000", date(2026, 9, 1))
    s = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert s["estimated_fuel_cost"] == "3.00"
    assert s["fuel_cost_status"] == "calculated"


async def test_sin_estado_operativo_espera_con_su_motivo(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-NS")
    await _jornada(alpha_client, seeded, millas="10.0")
    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))

    vista = await _informe(alpha_client, seeded, range="week")
    assert vista["summary"]["estimated_fuel_cost"] == "0.00"
    assert vista["summary"]["fuel_cost_status"] == "missing_input"
    assert vista["fuel"]["missing"] == {"operational_state": 1}


# ── Referencia de combustible ───────────────────────────────────────────────


async def test_la_formula_exacta(seeded, alpha_client):
    """10 millas / 25 MPG = 0,4 galones × 3,50 = 1,40."""
    perfil, vehiculo = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-F")
    await _vehiculo(alpha_client, seeded, vehiculo, operational_mpg="25.00")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _jornada(alpha_client, seeded, millas="10.0")
    await _precio(seeded.alpha.id, "GA", "regular", "3.5000", date(2026, 9, 1))

    s = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert _d(s["estimated_gallons"]) == Decimal("0.400")
    assert s["estimated_fuel_cost"] == "1.40"


async def test_la_referencia_aplicable_estado_grado_y_fecha(seeded, alpha_client):
    """La más reciente con fecha ≤ session_date; ni futura, ni de otro estado o grado."""
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-R")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _jornada(alpha_client, seeded, millas="24.0")  # 1 galón
    c = seeded.alpha.id
    await _precio(c, "GA", "regular", "2.0000", date(2026, 9, 1))
    await _precio(c, "GA", "regular", "3.0000", date(2026, 9, 10))  # la vigente
    await _precio(c, "GA", "regular", "9.0000", date(2026, 9, 17))  # futura
    await _precio(c, "TX", "regular", "8.0000", date(2026, 9, 15))  # otro estado
    await _precio(c, "GA", "premium", "7.0000", date(2026, 9, 15))  # otro grado

    s = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert s["estimated_fuel_cost"] == "3.00"


async def test_sin_referencia_espera_y_se_resuelve_cuando_llega(seeded, alpha_client):
    """Coste visible 0, estado distinto de calculado, sin congelar; luego, el real."""
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-W")
    await _estado(alpha_client, seeded, perfil, "GA")
    hecho = await _jornada(alpha_client, seeded, millas="24.0")
    await _precio(seeded.alpha.id, "TX", "regular", "8.0000", date(2026, 9, 1))  # no aplica

    espera = await _informe(alpha_client, seeded, range="week")
    assert espera["summary"]["estimated_fuel_cost"] == "0.00"
    assert espera["summary"]["fuel_cost_status"] == "waiting_reference"
    assert espera["fuel"]["waiting_reference"] == 1
    async with async_session_maker() as s:
        assert await s.scalar(
            text("SELECT count(*) FROM work_session_fuel_estimate WHERE work_session_id = :i"),
            {"i": hecho["jornada"]["id"]},
        ) == 0, "un cero por falta de precio no se congela"

    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))
    resuelta = await _informe(alpha_client, seeded, range="week")
    assert resuelta["summary"]["estimated_fuel_cost"] == "3.00"
    assert resuelta["summary"]["fuel_cost_status"] == "calculated"


async def test_un_cero_calculado_no_es_un_cero_en_espera(seeded, alpha_client):
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-Z")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _jornada(alpha_client, seeded, millas="0")
    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))

    s = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert s["estimated_fuel_cost"] == "0.00"
    assert s["fuel_cost_status"] == "calculated"


async def test_una_estimacion_final_no_cambia_con_precios_posteriores(seeded, alpha_client):
    """Se congela con su procedencia; una referencia cargada después no la mueve."""
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-H")
    await _estado(alpha_client, seeded, perfil, "GA")
    hecho = await _jornada(alpha_client, seeded, millas="24.0")
    ref = await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1), "Feed A")

    primera = await _informe(alpha_client, seeded, range="week")
    assert primera["summary"]["estimated_fuel_cost"] == "3.00"
    assert primera["fuel"]["final"] == 1

    async with async_session_maker() as s:
        e = (
            await s.execute(
                text(
                    "SELECT fuel_reference_id, state_code, fuel_grade, operational_mpg, "
                    "official_miles, price_per_gallon, effective_date, source, "
                    "estimated_gallons, estimated_cost, calculated_at "
                    "FROM work_session_fuel_estimate WHERE work_session_id = :i"
                ),
                {"i": hecho["jornada"]["id"]},
            )
        ).one()
    assert e.fuel_reference_id == ref and e.source == "Feed A"
    assert (e.state_code, e.fuel_grade) == ("GA", "regular")
    assert (e.operational_mpg, e.official_miles) == (Decimal("24.00"), Decimal("24.0000"))
    assert (e.estimated_gallons, e.estimated_cost) == (Decimal("1.0000"), Decimal("3.0000"))
    assert e.calculated_at is not None

    # Un precio nuevo, aplicable por fecha, llega después.
    await _precio(seeded.alpha.id, "GA", "regular", "5.0000", date(2026, 9, 15), "Feed B")
    despues = await _informe(alpha_client, seeded, range="week")
    assert despues["summary"]["estimated_fuel_cost"] == "3.00"


async def test_una_jornada_abierta_no_se_congela(seeded, alpha_client):
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-O")
    await _estado(alpha_client, seeded, perfil, "GA")
    hecho = await _jornada(alpha_client, seeded, millas="24.0", cerrar=False)
    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))

    vista = await _informe(alpha_client, seeded, range="week")
    assert vista["summary"]["estimated_fuel_cost"] == "3.00"
    assert vista["fuel"]["final"] == 0
    async with async_session_maker() as s:
        assert await s.scalar(
            text("SELECT count(*) FROM work_session_fuel_estimate WHERE work_session_id = :i"),
            {"i": hecho["jornada"]["id"]},
        ) == 0


async def test_mpg_distintos_se_calculan_por_jornada(seeded, alpha_client):
    """10 mi a 20 MPG + 10 mi a 40 MPG, a 4,00: 2,00 + 1,00 = 3,00 (no un promedio)."""
    perfil, vehiculo = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-MP")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _vehiculo(alpha_client, seeded, vehiculo, operational_mpg="20.00")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="10.0")
    await _vehiculo(alpha_client, seeded, vehiculo, operational_mpg="40.00")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 16), millas="10.0")
    await _precio(seeded.alpha.id, "GA", "regular", "4.0000", date(2026, 9, 1))

    s = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert s["estimated_fuel_cost"] == "3.00"


async def test_la_referencia_es_append_only_y_rechaza_precio_cero(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-AO")
    ref = await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))
    async with async_session_maker() as s:
        with pytest.raises(Exception, match="append-only"):
            await s.execute(
                text("UPDATE fuel_reference SET price_per_gallon = 1 WHERE id = :i"),
                {"i": ref},
            )
        await s.rollback()
    with pytest.raises(Exception, match="ck_fuel_reference_price_positive"):
        await _precio(seeded.alpha.id, "GA", "regular", "0", date(2026, 9, 2))


# ── Permisos y tenant ───────────────────────────────────────────────────────


async def test_permisos_de_reports_y_del_excel(seeded, alpha_client, alpha_anonymous):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-PR")
    params = {"range": "week", "date": MIERCOLES.isoformat()}

    assert (await alpha_anonymous.get(RUTA, params=params)).status_code == 401

    await alpha_client.login(seeded.alpha.users["supervisor"].email)
    assert (await alpha_client.get(RUTA, params=params)).status_code == 403
    assert (await alpha_client.get(f"{RUTA}/export", params=params)).status_code == 403

    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    assert (await alpha_client.get(RUTA, params=params)).status_code == 200
    assert (await alpha_client.get(f"{RUTA}/export", params=params)).status_code == 200

    # Leer sin exportar: el Excel exige las dos.
    async with async_session_maker() as s:
        await s.execute(
            text(
                "DELETE FROM role_permission WHERE role_id = :r AND permission_id = "
                "(SELECT id FROM permission WHERE name = 'route.reports.export')"
            ),
            {"r": seeded.alpha.roles["route_admin"]},
        )
        await s.commit()
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    assert (await alpha_client.get(RUTA, params=params)).status_code == 200
    assert (await alpha_client.get(f"{RUTA}/export", params=params)).status_code == 403


async def test_otro_tenant_y_un_id_inexistente_son_indistinguibles(
    seeded, alpha_client, beta_client
):
    await _perfil_y_vehiculo(beta_client, seeded, unidad="V-R9-BETA", compania="beta")
    await alpha_client.login(seeded.alpha.users["route_admin"].email)
    params = {"range": "week", "date": MIERCOLES.isoformat()}
    for ruta in (RUTA, f"{RUTA}/export"):
        ajeno = await alpha_client.get(
            ruta, params={**params, "supervisor_user_id": seeded.beta.users["supervisor"].id}
        )
        falso = await alpha_client.get(ruta, params={**params, "supervisor_user_id": 987654321})
        assert ajeno.status_code == falso.status_code == 404
        assert ajeno.json() == falso.json()


async def test_all_supervisors_no_incluye_otro_tenant(seeded, alpha_client, beta_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-TA")
    await _jornada(alpha_client, seeded, millas="10.0")
    await _perfil_y_vehiculo(beta_client, seeded, unidad="V-R9-TB", compania="beta")

    vista = await _informe(alpha_client, seeded, range="week")
    ids = {s["user_id"] for s in vista["supervisors"]}
    assert seeded.beta.users["supervisor"].id not in ids
    assert vista["summary"]["official_miles"] == "10.0"


async def test_la_alineacion_concede_reports_a_tenants_existentes(seeded):
    """Un tenant creado antes de RTE09 recibe las capacidades por alineación."""
    from app.db.scripts.align_role_capabilities import alinear

    consulta = text(
        "SELECT count(*) FROM role_permission rp JOIN permission p ON p.id = rp.permission_id "
        "WHERE rp.role_id = :r AND p.name IN ('route.reports.read', 'route.reports.export')"
    )
    async with async_session_maker() as s:
        await s.execute(
            text(
                "DELETE FROM role_permission WHERE role_id = :r AND permission_id IN "
                "(SELECT id FROM permission WHERE name LIKE 'route.reports.%')"
            ),
            {"r": seeded.alpha.roles["route_admin"]},
        )
        await s.commit()
        assert await s.scalar(consulta, {"r": seeded.alpha.roles["route_admin"]}) == 0

    await alinear()

    async with async_session_maker() as s:
        assert await s.scalar(consulta, {"r": seeded.alpha.roles["route_admin"]}) == 2
        assert await s.scalar(consulta, {"r": seeded.alpha.roles["supervisor"]}) == 0


async def test_el_estado_operativo_se_valida_audita_y_protege(
    seeded, alpha_client, beta_client
):
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-OS")
    ruta = f"/api/supervisors/{perfil['id']}/operational-state"

    assert (await alpha_client.put(ruta, json={"operational_state": "ZZ"})).status_code == 422
    assert (await alpha_client.put(ruta, json={"operational_state": "ga"})).json()[
        "operational_state"
    ] == "GA"

    await alpha_client.login(seeded.alpha.users["supervisor"].email)
    assert (await alpha_client.put(ruta, json={"operational_state": "TX"})).status_code == 403
    await beta_client.login(seeded.beta.users["route_admin"].email)
    assert (await beta_client.put(ruta, json={"operational_state": "TX"})).status_code == 404

    async with async_session_maker() as s:
        evento = (
            await s.execute(
                text(
                    "SELECT actor_user_id, changes FROM audit_event WHERE company_id = :c "
                    "AND action = 'operational_state'"
                ),
                {"c": seeded.alpha.id},
            )
        ).one()
    assert evento.changes == {"operational_state": {"old": None, "new": "GA"}}
    assert evento.actor_user_id == seeded.alpha.users["route_admin"].id


# ── Excel ───────────────────────────────────────────────────────────────────


async def _excel(cliente, seeded, **parametros) -> list[list]:
    await cliente.login(seeded.alpha.users["route_admin"].email)
    parametros.setdefault("date", MIERCOLES.isoformat())
    r = await cliente.get(f"{RUTA}/export", params=parametros)
    assert r.status_code == 200, r.text
    hoja = load_workbook(BytesIO(r.content))["Activity Detail"]
    return [list(fila) for fila in hoja.iter_rows(values_only=True)]


async def test_el_excel_tiene_las_16_columnas_y_concilia(seeded, alpha_client):
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-XL")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 15), millas="24.0", actividades=3)
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 16), millas="12.0")
    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 16))

    filas = await _excel(alpha_client, seeded, range="week")
    cabecera, datos = filas[0], filas[1:]
    assert cabecera == [
        "Date", "Supervisor", "Vehicle", "Activity Type", "Reference", "Start", "End",
        "Duration", "Trip Miles", "Operational MPG", "Fuel Grade", "Reference Fuel Price",
        "Estimated Gallons", "Estimated Fuel Cost", "Outcome", "Notes",
    ]
    assert len(datos) == 2, "una fila por viaje, aunque haya tres actividades"
    col = {nombre: i for i, nombre in enumerate(cabecera)}

    vista = await _informe(alpha_client, seeded, range="week")
    assert sum(Decimal(str(f[col["Trip Miles"]])) for f in datos) == _d(
        vista["summary"]["official_miles"]
    )
    assert sum(Decimal(str(f[col["Estimated Fuel Cost"]])) for f in datos) == _d(
        vista["summary"]["estimated_fuel_cost"]
    )

    del15, del16 = datos
    assert del15[col["Activity Type"]].startswith("Client Visit: ")
    assert del15[col["Activity Type"]].count(",") == 2
    # El 15 no tiene precio vigente (la referencia es del 16): vacío, coste 0.
    assert del15[col["Reference Fuel Price"]] is None
    assert del15[col["Estimated Fuel Cost"]] == 0
    assert del15[col["Estimated Gallons"]] == 1.0
    # El 16 sí.
    assert del16[col["Reference Fuel Price"]] == 3.0
    assert del16[col["Estimated Fuel Cost"]] == 1.5
    assert del16[col["Fuel Grade"]] == "Regular"


async def test_el_excel_respeta_supervisor_y_periodo(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-XS1")
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-XS2", usuario="route_admin")
    await _jornada(alpha_client, seeded, millas="10.0")
    await _jornada(alpha_client, seeded, millas="4.0", usuario="route_admin")
    await _jornada(alpha_client, seeded, dia=date(2026, 9, 2), millas="7.0")

    solo = await _excel(
        alpha_client, seeded, range="week",
        supervisor_user_id=seeded.alpha.users["route_admin"].id,
    )
    assert len(solo) == 2 and solo[1][8] == 4.0
    mes = await _excel(alpha_client, seeded, range="month")
    assert len(mes) == 4


async def test_un_valor_retirado_sigue_legible_en_el_excel(seeded, alpha_client):
    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-RV")
    hecho = await _jornada(alpha_client, seeded, millas="10.0", actividades=1)
    async with async_session_maker() as s:
        etiqueta = await s.scalar(
            text(
                "SELECT aea.label FROM activity_execution_activity aea "
                "JOIN activity_execution ae ON ae.id = aea.activity_execution_id "
                "WHERE ae.trip_id = :t"
            ),
            {"t": hecho["viaje"]["id"]},
        )
        await s.execute(
            text("UPDATE standard_value SET is_active = false WHERE company_id = :c"),
            {"c": seeded.alpha.id},
        )
        await s.commit()

    filas = await _excel(alpha_client, seeded, range="week")
    assert etiqueta in filas[1][3]


# ── Migración: el grado de las jornadas anteriores ──────────────────────────


async def test_el_relleno_del_grado_usa_el_de_entonces(seeded, alpha_client):
    """Una jornada anterior a RTE09 toma el grado que el vehículo tenía al empezar,
    aunque después se cambiara: lo reconstruye con la auditoría del vehículo."""
    from importlib import import_module

    relleno = import_module(
        "app.migrations.versions.0016_rte09_reports_fuel"
    ).RELLENO_GRADO_SQL

    _, vehiculo = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-BF")
    hecho = await _jornada(alpha_client, seeded, millas="10.0")
    await _vehiculo(alpha_client, seeded, vehiculo, fuel_grade="premium")
    await _vehiculo(alpha_client, seeded, vehiculo, fuel_grade="diesel")

    async with async_session_maker() as s:
        # Como si la jornada fuera anterior a 0016: sin grado congelado.
        await s.execute(
            text("UPDATE work_session SET fuel_grade_snapshot = NULL WHERE id = :i"),
            {"i": hecho["jornada"]["id"]},
        )
        await s.execute(text(relleno))
        await s.commit()
        grado = await s.scalar(
            text("SELECT fuel_grade_snapshot FROM work_session WHERE id = :i"),
            {"i": hecho["jornada"]["id"]},
        )
    assert grado == "regular", "el de entonces, no el actual (diesel)"


# ── Evidencia del entregable (§25): las dos cadenas, enlace a enlace ────────


def _imprimir(titulo: str, cadena: list[tuple[str, object]]) -> None:
    print(f"\n{titulo}")
    for enlace, valor in cadena:
        print(f"  → {enlace}: {valor}")


async def _snapshots(jornada_id: int):
    async with async_session_maker() as s:
        return (
            await s.execute(
                text(
                    "SELECT session_date, operational_state_snapshot, fuel_grade_snapshot, "
                    "mpg_snapshot, vehicle_id FROM work_session WHERE id = :i"
                ),
                {"i": jornada_id},
            )
        ).one()


async def test_evidencia_cadena_completa_calculada(seeded, alpha_client):
    """Supervisor → estado → jornada → snapshots → viajes → millas → referencia →
    estimación → API → Excel. Dos viajes, uno con dos actividades."""
    perfil, vehiculo = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-EV1")
    await _estado(alpha_client, seeded, perfil, "GA")
    a = await _jornada_con_parada(alpha_client, seeded, actividades=2, cerrar_jornada=False)
    await _millas(a["viaje"]["id"], "18.0")
    await alpha_client.login(seeded.alpha.users["supervisor"].email)
    oficina = (await _valores(alpha_client, seeded.alpha.id, "office_purposes"))[0]["id"]
    plan = {"purpose": "office", "context_reference": "Main Office", "standard_value_id": oficina}
    b = (await alpha_client.post("/api/trips", json=plan)).json()
    assert "id" in b, b
    await alpha_client.post(f"/api/trips/{b['id']}/start", json={})
    await alpha_client.post(f"/api/trips/{b['id']}/arrive", json={})
    await alpha_client.post(f"/api/trips/{b['id']}/activity/start", json={"activity_ids": []})
    resultado = (await _valores(alpha_client, seeded.alpha.id, "outcomes"))[0]["id"]
    fin = await alpha_client.post(
        f"/api/trips/{b['id']}/activity/complete",
        json={"action": "complete", "outcome_id": resultado},
    )
    assert fin.status_code == 200, fin.text
    await _millas(b["id"], "6.0")
    jornada = a["jornada"]["id"]
    # El cierre pasa por su guarda de odómetro, como en el producto.
    await alpha_client.post(
        f"/api/odometer/sessions/{jornada}/end/photo",
        files={"photo": ("odo.png", FOTO, "image/png")},
    )
    await alpha_client.post(
        f"/api/odometer/sessions/{jornada}/end/confirm", json={"reading": "100120.0"}
    )
    cierre = await alpha_client.post(f"/api/worksessions/{jornada}/end", json={})
    assert cierre.status_code == 200, cierre.text
    await _fechar_jornada(jornada, MIERCOLES)
    # Después de la jornada: el supervisor y el vehículo cambian; no la historia.
    await _estado(alpha_client, seeded, perfil, "TX")
    await _vehiculo(alpha_client, seeded, vehiculo, fuel_grade="diesel", operational_mpg="10.00")
    ref = await _precio(seeded.alpha.id, "GA", "regular", "3.1250", date(2026, 9, 14), "Evidence feed")
    await _precio(seeded.alpha.id, "TX", "diesel", "4.0000", date(2026, 9, 14))

    snap = await _snapshots(jornada)
    vista = await _informe(alpha_client, seeded, range="week")
    filas = await _excel(alpha_client, seeded, range="week")
    async with async_session_maker() as s:
        congelada = (
            await s.execute(
                text(
                    "SELECT fuel_reference_id, price_per_gallon, effective_date, source, "
                    "estimated_gallons, estimated_cost FROM work_session_fuel_estimate "
                    "WHERE work_session_id = :i"
                ),
                {"i": jornada},
            )
        ).one()

    assert (snap.operational_state_snapshot, snap.fuel_grade_snapshot) == ("GA", "regular")
    assert snap.mpg_snapshot == Decimal("24.00") and snap.session_date == MIERCOLES
    s = vista["summary"]
    assert s["official_miles"] == "24.0" and s["activities"] == 2
    assert _d(s["estimated_gallons"]) == Decimal("1.000")
    assert s["estimated_fuel_cost"] == "3.13" and s["fuel_cost_status"] == "calculated"
    assert congelada.fuel_reference_id == ref and congelada.estimated_cost == Decimal("3.1250")
    assert len(filas) == 3, "dos viajes, dos filas; las dos actividades van en una"
    assert sum(f[13] for f in filas[1:]) == pytest.approx(3.125, abs=0.01)
    assert {f[11] for f in filas[1:]} == {3.125}

    _imprimir("EVIDENCIA 1 · cadena calculada", [
        ("Supervisor", seeded.alpha.users["supervisor"].email),
        ("Operational State al empezar (hoy TX)", snap.operational_state_snapshot),
        ("Work Session / session_date", f"#{jornada} / {snap.session_date}"),
        ("Vehicle / MPG snapshot (hoy 10.00)", f"#{snap.vehicle_id} / {snap.mpg_snapshot}"),
        ("Fuel Grade snapshot (hoy diesel)", snap.fuel_grade_snapshot),
        ("Trips", f"#{a['viaje']['id']} client_visit 18.0 mi, 2 actividades; "
                  f"#{b['id']} office 6.0 mi"),
        ("Official Miles", s["official_miles"]),
        ("Fuel Reference", f"#{ref} GA/regular 3.1250 desde 2026-09-14, 'Evidence feed' "
                           "(ignorada la TX/diesel actual)"),
        ("Estado", s["fuel_cost_status"]),
        ("Estimated Gallons", s["estimated_gallons"]),
        ("Est. Fuel (API)", s["estimated_fuel_cost"]),
        ("Congelada", f"ref #{congelada.fuel_reference_id} · {congelada.price_per_gallon} · "
                      f"{congelada.effective_date} · {congelada.source} · "
                      f"{congelada.estimated_gallons} gal · {congelada.estimated_cost}"),
        ("Breakdown", [(r["purpose"], r["official_miles"], r["estimated_fuel_cost"])
                       for r in vista["breakdown"]]),
        ("Excel", [(f[3], f[8], f[11], f[12], f[13]) for f in filas[1:]]),
    ])


async def test_evidencia_sin_referencia(seeded, alpha_client):
    """Sin precio: 0 visible, estado ≠ calculado, ningún precio inventado, y se
    resuelve cuando llega la referencia."""
    perfil, _ = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-EV2")
    await _estado(alpha_client, seeded, perfil, "GA")
    hecho = await _jornada(alpha_client, seeded, millas="24.0")
    jornada = hecho["jornada"]["id"]

    vista = await _informe(alpha_client, seeded, range="week")
    filas = await _excel(alpha_client, seeded, range="week")
    async with async_session_maker() as s:
        precios = await s.scalar(text("SELECT count(*) FROM fuel_reference"))
        congeladas = await s.scalar(
            text("SELECT count(*) FROM work_session_fuel_estimate WHERE work_session_id = :i"),
            {"i": jornada},
        )
    s0 = vista["summary"]
    assert s0["estimated_fuel_cost"] == "0.00"
    assert s0["fuel_cost_status"] == "waiting_reference" != "calculated"
    assert precios == 0 and congeladas == 0
    assert filas[1][11] is None and filas[1][13] == 0

    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1), "Late feed")
    despues = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert despues["estimated_fuel_cost"] == "3.00"
    assert despues["fuel_cost_status"] == "calculated"

    _imprimir("EVIDENCIA 2 · sin referencia", [
        ("Est. Fuel visible", s0["estimated_fuel_cost"]),
        ("Estado interno", s0["fuel_cost_status"]),
        ("Filas en fuel_reference", precios),
        ("Estimaciones congeladas", congeladas),
        ("Excel · Reference Fuel Price / Est. Fuel Cost", (filas[1][11], filas[1][13])),
        ("Llega la referencia GA/regular 3.0000", "→"),
        ("Est. Fuel / estado", (despues["estimated_fuel_cost"], despues["fuel_cost_status"])),
    ])


async def test_el_grado_congelado_elige_su_precio(seeded, alpha_client):
    """Diésel al empezar, regular después: aplica el precio del diésel."""
    perfil, vehiculo = await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-GD")
    await _estado(alpha_client, seeded, perfil, "GA")
    await _vehiculo(alpha_client, seeded, vehiculo, fuel_grade="diesel")
    await _jornada(alpha_client, seeded, millas="24.0")  # 1 galón
    await _vehiculo(alpha_client, seeded, vehiculo, fuel_grade="regular")
    await _precio(seeded.alpha.id, "GA", "diesel", "4.0000", date(2026, 9, 1))
    await _precio(seeded.alpha.id, "GA", "regular", "3.0000", date(2026, 9, 1))

    s = (await _informe(alpha_client, seeded, range="week"))["summary"]
    assert s["estimated_fuel_cost"] == "4.00"
    filas = await _excel(alpha_client, seeded, range="week")
    assert filas[1][10] == "Diesel"


async def test_la_consulta_filtra_por_compania_por_si_misma(seeded, alpha_client):
    """Aunque alguien le pase ids de otra compañía, la consulta no los cruza:
    el alcance no depende sólo de la validación del router."""
    from app.routers_api.reports.dao import hechos

    await _perfil_y_vehiculo(alpha_client, seeded, unidad="V-R9-DAO")
    await _jornada(alpha_client, seeded, millas="10.0")
    supervisor = seeded.alpha.users["supervisor"].id
    rango = {"inicio": date(2026, 9, 14), "fin": date(2026, 9, 20)}

    assert len(await hechos(company_id=seeded.alpha.id, user_ids=[supervisor], **rango)) == 1
    assert await hechos(company_id=seeded.beta.id, user_ids=[supervisor], **rango) == []
